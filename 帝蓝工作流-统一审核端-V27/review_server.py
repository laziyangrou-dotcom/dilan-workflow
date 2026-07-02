# -*- coding: utf-8 -*-
"""帝蓝工作流 - 统一审核端
支持素材 / 图片 / 视频工程包的本地导入、审核、制作流程预览、表单导出和审核结果工程包导出。
"""
import copy
import datetime as _dt
import io
import json
import mimetypes
import os
import re
import shutil
import sys
import tempfile
import threading
import time
import uuid
import webbrowser
import zipfile
from pathlib import Path
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import quote, unquote, urlparse, parse_qs

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except Exception:
    openpyxl = None

ROOT = Path(__file__).resolve().parent
PORT = int(os.environ.get('REVIEW_CLIENT_PORT', '8797'))
WORKSPACES = ROOT / 'review_workspace'
WORKSPACES.mkdir(parents=True, exist_ok=True)
REVIEW_PARENTS_JSON = WORKSPACES / 'parents.json'
PROJECT_MODULE_KEYS = ['material', 'image', 'video']
PROJECT_SEP = '__CHILD__'

IMAGE_EXTS = {'.png','.jpg','.jpeg','.webp','.gif','.bmp'}
VIDEO_EXTS = {'.mp4','.mov','.webm','.mkv','.avi'}
MODULE_LABELS = {'material':'美术模块','image':'分镜模块','video':'视频模块'}
MODULE_UI_LABELS = {'material':'美术','image':'分镜','video':'视频'}
import sys as _sys
if str(ROOT) not in _sys.path:
    _sys.path.insert(0, str(ROOT))
import datacenter_common as dcc
CATEGORIES = ['人物','场景','道具']
GROUPED_CATEGORIES = {'人物','场景'}


def now_str():
    return _dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')

def stamp():
    return _dt.datetime.now().strftime('%Y%m%d_%H%M%S')

def new_id(prefix='id'):
    return f"{prefix}_{uuid.uuid4().hex[:12]}"

def safe_name(s, default='item'):
    s = str(s or default).strip()
    s = re.sub(r'[\\/:*?"<>|\r\n\t]+','_',s)
    s = re.sub(r'\s+','_',s)
    return s[:120] or default

def read_json(path, default=None):
    try:
        return json.loads(Path(path).read_text(encoding='utf-8'))
    except Exception:
        return copy.deepcopy(default)

def write_json(path, data):
    p=Path(path); p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')


def split_project_ref(project_id):
    raw = str(project_id or '').strip()
    if PROJECT_SEP in raw:
        parent, child = raw.split(PROJECT_SEP, 1)
        return safe_name(parent, ''), safe_name(child, '')
    return '', safe_name(raw, '')

def load_review_parent_index():
    data = read_json(REVIEW_PARENTS_JSON, None)
    if not isinstance(data, dict):
        data = {'version': 1, 'parents': []}
    if not isinstance(data.get('parents'), list):
        data['parents'] = []
    return data

def save_review_parent_index(data):
    data = data if isinstance(data, dict) else {'version': 1, 'parents': []}
    data['version'] = data.get('version') or 1
    data['parents'] = data.get('parents') if isinstance(data.get('parents'), list) else []
    write_json(REVIEW_PARENTS_JSON, data)

def merge_review_parent(parent_id, parent_name=None, created_at=None, updated_at=None):
    pid = safe_name(parent_id, '导入项目')
    pname = safe_name(parent_name or pid, pid)
    data = load_review_parent_index()
    now = now_str()
    found = None
    for item in data.get('parents', []):
        if safe_name(item.get('parent_id') or '', '') == pid:
            found = item
            break
    if not found:
        found = {'parent_id': pid, 'parent_name': pname, 'created_at': created_at or now, 'updated_at': updated_at or now}
        data['parents'].append(found)
    else:
        found['parent_id'] = pid
        found['parent_name'] = pname or found.get('parent_name') or pid
        found['created_at'] = found.get('created_at') or created_at or now
        found['updated_at'] = updated_at or now
    save_review_parent_index(data)
    return found

def remove_review_parent(parent_id):
    pid = safe_name(parent_id, '')
    data = load_review_parent_index()
    data['parents'] = [p for p in data.get('parents', []) if safe_name(p.get('parent_id') or '', '') != pid]
    save_review_parent_index(data)

def parent_info_from_workspace(info, project=None, manifest=None):
    project = project if isinstance(project, dict) else {}
    manifest = manifest if isinstance(manifest, dict) else {}
    scope = manifest.get('project_scope') if isinstance(manifest.get('project_scope'), dict) else {}
    old_parent, old_child = split_project_ref(project.get('project_id') or manifest.get('source_project_id') or '')
    parent_id = info.get('parent_id') or project.get('parent_id') or manifest.get('parent_id') or scope.get('parent_id') or old_parent or '未归类项目'
    parent_name = info.get('parent_name') or project.get('parent_name') or manifest.get('parent_name') or scope.get('parent_name') or parent_id
    child_id = info.get('child_id') or project.get('child_id') or manifest.get('child_id') or scope.get('child_id') or old_child or project.get('project_id') or '未命名子项目'
    child_name = info.get('child_name') or project.get('child_name') or manifest.get('child_name') or scope.get('child_name') or project.get('project_name') or child_id
    return {
        'parent_id': safe_name(parent_id, '未归类项目'),
        'parent_name': safe_name(parent_name or parent_id, '未归类项目'),
        'child_id': safe_name(child_id, '未命名子项目'),
        'child_name': safe_name(child_name or child_id, '未命名子项目'),
    }

def collect_review_parent_items():
    parents = {}
    data = load_review_parent_index()
    for item in data.get('parents', []):
        pid = safe_name(item.get('parent_id') or '', '')
        if not pid:
            continue
        parents[pid] = {
            'parent_id': pid,
            'parent_name': item.get('parent_name') or pid,
            'created_at': item.get('created_at') or '',
            'updated_at': item.get('updated_at') or '',
            'module_counts': {m: 0 for m in PROJECT_MODULE_KEYS},
            'child_count': 0,
            'total_child_count': 0,
        }
    if WORKSPACES.exists():
        for w in WORKSPACES.iterdir():
            if not w.is_dir():
                continue
            info = read_json(w/'workspace.json', {}) or {}
            project = read_json(w/'project.json', {}) or {}
            manifest = read_json(w/'manifest.json', {}) or {}
            if not project:
                continue
            module_type = info.get('module_type') or detect_module(manifest, project, info.get('project_json_rel', ''))
            pinf = parent_info_from_workspace(info, project, manifest)
            pid = pinf['parent_id']
            item = parents.setdefault(pid, {
                'parent_id': pid,
                'parent_name': pinf['parent_name'],
                'created_at': info.get('imported_at') or '',
                'updated_at': info.get('imported_at') or '',
                'module_counts': {m: 0 for m in PROJECT_MODULE_KEYS},
                'child_count': 0,
                'total_child_count': 0,
            })
            if pinf.get('parent_name') and (not item.get('parent_name') or item.get('parent_name') == pid):
                item['parent_name'] = pinf['parent_name']
            ts = str(info.get('imported_at') or '')
            if ts:
                if not item.get('created_at'):
                    item['created_at'] = ts
                if ts > str(item.get('updated_at') or ''):
                    item['updated_at'] = ts
            if module_type in PROJECT_MODULE_KEYS:
                item['module_counts'][module_type] = int(item['module_counts'].get(module_type) or 0) + 1
                item['total_child_count'] += 1
    for item in parents.values():
        item['child_count'] = item.get('total_child_count') or 0
        item['module_labels'] = MODULE_UI_LABELS
    merged = {'version': 1, 'parents': [{k: v for k, v in item.items() if k in ('parent_id', 'parent_name', 'created_at', 'updated_at')} for item in parents.values()]}
    save_review_parent_index(merged)
    return sorted(parents.values(), key=lambda x: x.get('updated_at') or '', reverse=True)

def find_review_parent(parent_id=None, parent_name=None):
    pid = safe_name(parent_id or '', '')
    pname = safe_name(parent_name or '', '')
    if not pid and not pname:
        return None
    parents = collect_review_parent_items()
    for item in parents:
        if pid and safe_name(item.get('parent_id') or '', '') == pid:
            return {'parent_id': item['parent_id'], 'parent_name': item.get('parent_name') or item['parent_id'], 'matched_by': 'parent_id'}
    for item in parents:
        if pname and safe_name(item.get('parent_name') or item.get('parent_id') or '', '') == pname:
            return {'parent_id': item['parent_id'], 'parent_name': item.get('parent_name') or item['parent_id'], 'matched_by': 'parent_name'}
    return None

def read_parent_json_from_extract(extracted, manifest):
    rel = (manifest or {}).get('parent_json')
    if rel:
        p = Path(extracted) / str(rel).lstrip('/')
        if p.exists():
            data = read_json(p, {}) or {}
            if isinstance(data, dict):
                return data
    for p in Path(extracted).rglob('parent.json'):
        data = read_json(p, {}) or {}
        if isinstance(data, dict) and (data.get('parent_id') or data.get('parent_name')):
            return data
    return {}

def resolve_review_import_scope(manifest, project, parent_data, filename=''):
    manifest = manifest if isinstance(manifest, dict) else {}
    project = project if isinstance(project, dict) else {}
    parent_data = parent_data if isinstance(parent_data, dict) else {}
    scope = manifest.get('project_scope') if isinstance(manifest.get('project_scope'), dict) else {}
    old_parent, old_child = split_project_ref(project.get('project_id') or manifest.get('source_project_id') or manifest.get('project_id') or '')
    source_parent_id = safe_name(
        project.get('parent_id') or manifest.get('parent_id') or scope.get('parent_id') or parent_data.get('parent_id') or old_parent or '导入项目',
        '导入项目'
    )
    source_parent_name = safe_name(
        project.get('parent_name') or manifest.get('parent_name') or scope.get('parent_name') or parent_data.get('parent_name') or source_parent_id,
        source_parent_id
    )
    matched = find_review_parent(source_parent_id, source_parent_name)
    if matched:
        target_parent_id = matched['parent_id']
        target_parent_name = matched.get('parent_name') or target_parent_id
        matched_by = matched.get('matched_by') or 'existing'
        created_new_parent = False
    else:
        target_parent_id = source_parent_id
        target_parent_name = source_parent_name
        matched_by = 'created_from_package'
        created_new_parent = True
    source_child_id = safe_name(
        project.get('child_id') or manifest.get('child_id') or scope.get('child_id') or old_child or project.get('project_id') or project.get('project_name') or '导入子项目',
        '导入子项目'
    )
    source_child_name = safe_name(
        project.get('child_name') or manifest.get('child_name') or scope.get('child_name') or project.get('project_name') or manifest.get('project_name') or source_child_id,
        source_child_id
    )
    return {
        'source_parent_id': source_parent_id,
        'source_parent_name': source_parent_name,
        'parent_id': target_parent_id,
        'parent_name': target_parent_name,
        'matched_by': matched_by,
        'created_new_parent': created_new_parent,
        'source_child_id': source_child_id,
        'source_child_name': source_child_name,
    }

def choose_review_child_identity(parent_id, module_type, child_id, child_name):
    pid = safe_name(parent_id, '导入项目')
    cid_base = safe_name(child_id or child_name, '导入子项目')
    cname_base = safe_name(child_name or child_id, cid_base)
    existing_ids = set()
    existing_names = set()
    for item in list_workspace_infos():
        if item.get('parent_id') == pid and item.get('module_type') == module_type:
            existing_ids.add(safe_name(item.get('child_id') or '', ''))
            existing_names.add(safe_name(item.get('child_name') or item.get('project_name') or '', ''))
    if cid_base not in existing_ids and cname_base not in existing_names:
        return cid_base, cname_base
    idx = 1
    while True:
        cid = safe_name(f'{cid_base}_导入{idx}', '导入子项目')
        cname = safe_name(f'{cname_base}（导入{idx}）', '导入子项目')
        if cid not in existing_ids and cname not in existing_names:
            return cid, cname
        idx += 1

def parse_multipart_upload(headers, rfile):
    ctype=headers.get('Content-Type') or ''
    m=re.search(r'boundary=(.+)',ctype)
    if not m: raise ValueError('缺少 multipart boundary')
    boundary=m.group(1).strip().strip('"')
    length=int(headers.get('Content-Length') or 0)
    body=rfile.read(length) if length else b''
    marker=('--'+boundary).encode('utf-8')
    for part in body.split(marker):
        part=part.strip(b'\r\n')
        if not part or part==b'--': continue
        if b'\r\n\r\n' not in part: continue
        head,data=part.split(b'\r\n\r\n',1)
        if data.endswith(b'\r\n'): data=data[:-2]
        if data.endswith(b'--'): data=data[:-2]
        text=head.decode('utf-8',errors='ignore')
        if 'name="file"' not in text: continue
        fm=re.search(r'filename="([^"]*)"', text)
        return (fm.group(1) if fm else 'project.zip'), data
    raise ValueError('没有找到上传文件')

def safe_zip_members(zf):
    for info in zf.infolist():
        name=info.filename.replace('\\','/')
        if not name or name.endswith('/'): continue
        if name.startswith('/') or name.startswith('../') or '/../' in name: continue
        yield info,name

def safe_extract_zip(zip_path, dest):
    dest=Path(dest)
    with zipfile.ZipFile(zip_path,'r') as zf:
        for info,name in safe_zip_members(zf):
            target=dest/name
            target.parent.mkdir(parents=True,exist_ok=True)
            with zf.open(info) as src, open(target,'wb') as out:
                shutil.copyfileobj(src,out)

def load_manifest(extract_root):
    p=Path(extract_root)/'manifest.json'
    return read_json(p,{}) if p.exists() else {}

def find_project_json(extract_root, manifest=None):
    root=Path(extract_root)
    manifest=manifest or {}
    candidates=[]
    if manifest.get('project_json'):
        candidates.append(str(manifest.get('project_json')).lstrip('/'))
    for c in candidates:
        p=root/c
        if p.exists(): return p,c
    allp=list(root.rglob('project.json'))
    if not allp: raise FileNotFoundError('工程包中没有找到 project.json')
    p=allp[0]
    return p,p.relative_to(root).as_posix()

def file_records_paths(manifest):
    out=[]
    for k in ('project_json','parent_json','input_root','output_image_root','output_video_root','output_material_root','temp_refs_root','review_notes_root'):
        v=manifest.get(k)
        if v: out.append(str(v))
    for rec in manifest.get('files') or []:
        if isinstance(rec,dict) and rec.get('path'):
            out.append(str(rec.get('path')))
    return out

def detect_module(manifest, project, project_json_rel=''):
    """识别工程包所属模块。优先使用显式字段，旧包再走启发式。"""
    manifest = manifest or {}
    project = project or {}
    m = manifest.get('module_type') or project.get('module_type') or project.get('tool_module') or project.get('source_module')
    if m in ('material','image','video'):
        return m
    pkg = str(manifest.get('package_type') or '').lower()
    if 'asset' in pkg and 'package' in pkg:
        return 'material'
    paths = '\n'.join(file_records_paths(manifest)+[project_json_rel]).lower().replace('\\','/')
    if 'output/video/' in paths or 'output_videos' in paths or 'video_root' in paths or any_video_in_project(project):
        return 'video'
    if 'output/material/' in paths or 'output_material' in paths or '/material/' in paths or 'tools/material' in paths:
        return 'material'
    settings = project.get('project_settings') or {}
    settings_text = json.dumps(settings, ensure_ascii=False).lower()
    if 'output/material' in settings_text or '美术模块' in settings_text or '素材模块' in settings_text or 'material' in settings_text:
        return 'material'
    assets = project.get('assets') if isinstance(project.get('assets'), list) else []
    groups = project.get('asset_groups') if isinstance(project.get('asset_groups'), list) else []
    scenes = project.get('scenes') if isinstance(project.get('scenes'), list) else []
    candidate_count = 0
    generated_image_count = 0
    generated_video_count = 0
    for sc in scenes:
        for sh in sc.get('shots') or []:
            candidate_count += len(sh.get('storyboard_candidates') or [])
            for t in sh.get('tabs') or []:
                generated_image_count += len(t.get('generated_images') or [])
                generated_video_count += len(t.get('generated_videos') or []) + len(t.get('videos') or [])
    if generated_video_count > 0:
        return 'video'
    if (assets or groups) and candidate_count == 0 and generated_image_count == 0 and generated_video_count == 0:
        return 'material'
    group_cats = {g.get('category') for g in groups if isinstance(g, dict)}
    if groups and group_cats.intersection({'人物','场景','道具','物品'}) and candidate_count == 0:
        return 'material'
    return 'image'


def any_video_in_project(project):
    for sc in project.get('scenes') or []:
        for sh in sc.get('shots') or []:
            for c in sh.get('storyboard_candidates') or []:
                if is_video_path(image_path_of(c)): return True
            for t in sh.get('tabs') or []:
                for v in t.get('generated_videos') or []:
                    if is_video_path(image_path_of(v)): return True
                for v in t.get('videos') or []:
                    if is_video_path(image_path_of(v)): return True
    return False

def workspace_dir(wid):
    return WORKSPACES / safe_name(wid,'workspace')

def load_workspace(wid):
    w=workspace_dir(wid)
    info=read_json(w/'workspace.json',{}) or {}
    project=read_json(w/'project.json',{}) or {}
    manifest=read_json(w/'manifest.json',{}) or {}
    usage=read_json(w/'usage_summary.json',{}) or {}
    return w,info,project,manifest,usage

def save_workspace_project(wid, project):
    w=workspace_dir(wid)
    write_json(w/'project.json',project)


def list_workspace_infos():
    items=[]
    if not WORKSPACES.exists():
        return items
    for w in sorted(WORKSPACES.iterdir(), key=lambda p: p.stat().st_mtime if p.exists() else 0, reverse=True):
        if not w.is_dir():
            continue
        wid=w.name
        info=read_json(w/'workspace.json',{}) or {}
        project=read_json(w/'project.json',{}) or {}
        manifest=read_json(w/'manifest.json',{}) or {}
        usage=read_json(w/'usage_summary.json',{}) or {}
        if not project:
            continue
        module_type=info.get('module_type') or detect_module(manifest,project,info.get('project_json_rel',''))
        pinf=parent_info_from_workspace(info,project,manifest)
        stats=project_stats(project,module_type)
        usage_total=(usage.get('total') or {}) if isinstance(usage,dict) else {}
        items.append({
            'workspace': wid,
            'parent_id': pinf['parent_id'],
            'parent_name': pinf['parent_name'],
            'child_id': pinf['child_id'],
            'child_name': pinf['child_name'],
            'module_type': module_type,
            'module_label': MODULE_LABELS.get(module_type,module_type),
            'module_ui_label': MODULE_UI_LABELS.get(module_type,module_type),
            'project_name': pinf['child_name'] or project.get('project_name') or project.get('project_id') or info.get('source_filename') or '未命名工程',
            'project_id': project.get('project_id') or '',
            'source_filename': info.get('source_filename') or '',
            'imported_at': info.get('imported_at') or '',
            'exported_at': manifest.get('exported_at') or (project.get('export_owner') or {}).get('exported_at') or '',
            'responsible_nickname': manifest.get('responsible_nickname') or (project.get('export_owner') or {}).get('responsible_nickname') or '',
            'stats': stats,
            'usage_usd': round(float(usage_total.get('usd') or 0),6),
            'usage_tokens': int(usage_total.get('tokens') or 0),
        })
    return items

def usage_total_usd(usage):
    try:
        return round(float(((usage or {}).get('total') or {}).get('usd') or 0), 6)
    except Exception:
        return 0.0

def usage_scene_usd_map(usage):
    out={}
    for item in (usage or {}).get('by_sc') or []:
        code=str(item.get('scene_code') or '').strip()
        if code:
            out[code]=round(float(item.get('usd') or 0),6)
    return out

def first_time_from_items(items):
    keys=('created_at','generated_at','time','updated_at','submitted_at','finished_at')
    for it in items or []:
        if not isinstance(it,dict):
            continue
        for k in keys:
            v=it.get(k)
            if v:
                return str(v)
    return ''

def shot_generation_time(shot):
    items=[]
    items.extend(shot.get('storyboard_candidates') or [])
    for tab in shot.get('tabs') or []:
        items.extend(tab.get('generated_images') or [])
        items.extend(tab.get('generated_videos') or [])
        items.extend(tab.get('videos') or [])
        items.extend(tab.get('messages') or [])
    return first_time_from_items(items)

def parse_positive_int(v, default=None):
    try:
        if v in (None,''): return default
        n=int(float(v))
        return n if n>0 else default
    except Exception:
        return default

def scene_number(scene):
    n=parse_positive_int(scene.get('scene_number'),None)
    if n is not None: return n
    m=re.search(r'(\d+)', str(scene.get('scene_code') or ''))
    return int(m.group(1)) if m else 999999

def scene_code(scene):
    n=scene_number(scene)
    return str(scene.get('scene_code') or 'SC-0000') if n==999999 else f'SC-{n:04d}'

def review_label(n):
    return {1:'一审',2:'二审',3:'三审'}.get(int(n or 1),'一审')

def ensure_rounds(review):
    review.setdefault('current_round',1)
    review.setdefault('max_round_available',1)
    review.setdefault('rounds',{})
    if '1' not in review['rounds']:
        review['rounds']['1']={'label':'一审','status':'pending','locked':False,'submitted_at':'','submitted_by':'','note':''}
    maxn=1
    for k,r in list(review['rounds'].items()):
        n=parse_positive_int(k,1) or 1
        maxn=max(maxn,n)
        r.setdefault('label',review_label(n)); r.setdefault('status','pending'); r.setdefault('locked',False)
        r.setdefault('submitted_at',''); r.setdefault('submitted_by',''); r.setdefault('note','')
    review['max_round_available']=max(int(review.get('max_round_available') or 1), maxn)
    review['current_round']=min(max(int(review.get('current_round') or 1),1), int(review.get('max_round_available') or 1))
    return review

def ensure_scene_review(scene):
    rv=scene.setdefault('review',{})
    ensure_rounds(rv)
    # 兼容旧审核端：轮次里保留 shots 快照，但新逻辑不强依赖。
    for k,r in rv.get('rounds',{}).items():
        r.setdefault('shots', copy.deepcopy(scene.get('shots') or []))
    return rv

def material_unit_id(kind, raw_id):
    return f'{kind}:{raw_id}'

def ensure_material_review(project):
    mr=project.setdefault('material_review',{})
    mr.setdefault('current_category','人物')
    mr.setdefault('units',{})
    units=mr['units']
    for unit in material_units(project):
        uid=unit['unit_id']
        units.setdefault(uid, {'current_round':1,'max_round_available':1,'rounds':{}})
        ensure_rounds(units[uid])
    return mr

DOC_ASSET_EXTS={'.pdf','.doc','.docx','.txt','.md','.csv','.xls','.xlsx'}
def is_document_asset_rv(a):
    if not isinstance(a,dict): return False
    if a.get('asset_type')=='document': return True
    p=str((a.get('file_path') or a.get('path') or '')).split('?')[0]
    return Path(p).suffix.lower() in DOC_ASSET_EXTS

def normalize_assets(project):
    assets=project.get('assets') if isinstance(project.get('assets'),list) else []
    groups=project.get('asset_groups') if isinstance(project.get('asset_groups'),list) else []
    for a in assets:
        if not isinstance(a,dict): continue
        a.setdefault('asset_id', new_id('asset'))
        cat=a.get('category') if a.get('category') in CATEGORIES else ('道具' if a.get('category')=='物品' else '道具')
        if cat=='物品': cat='道具'
        a['category']=cat
        a.setdefault('name', Path(str(asset_path(a) or '素材')).stem or '素材')
    for g in groups:
        if not isinstance(g,dict): continue
        g.setdefault('group_id', new_id('grp'))
        if g.get('category')=='物品': g['category']='道具'
        g.setdefault('category','人物')
        g.setdefault('name','未命名'+str(g.get('category') or '素材'))
    project['assets']=assets; project['asset_groups']=groups

def material_units(project):
    normalize_assets(project)
    assets=[a for a in (project.get('assets') or []) if not is_document_asset_rv(a)]
    groups=project.get('asset_groups') or []
    out=[]
    for cat in ['人物','场景']:
        cat_groups=[g for g in groups if g.get('category')==cat]
        # 没有组但有素材时，生成一个虚拟未分组。
        grouped_ids={g.get('group_id') for g in cat_groups}
        loose=[a for a in assets if a.get('category')==cat and not a.get('temporary') and a.get('group_id') not in grouped_ids]
        for g in sorted(cat_groups, key=lambda x:int(x.get('sort_order') or 999999)):
            items=[a for a in assets if a.get('category')==cat and not a.get('temporary') and a.get('group_id')==g.get('group_id')]
            out.append({'unit_id':material_unit_id('group',g.get('group_id')),'unit_type':'group','category':cat,'title':g.get('name') or cat,'group':g,'assets':items})
        if loose:
            out.append({'unit_id':material_unit_id('group',f'loose_{cat}'),'unit_type':'group','category':cat,'title':'未分组'+cat,'group':{'group_id':f'loose_{cat}','name':'未分组'+cat,'category':cat},'assets':loose})
    for a in assets:
        if a.get('category') in ('道具','物品') and not a.get('temporary'):
            out.append({'unit_id':material_unit_id('asset',a.get('asset_id')),'unit_type':'asset','category':'道具','title':a.get('name') or '道具','asset':a,'assets':[a]})
    return out

def find_material_unit(project, unit_id):
    for u in material_units(project):
        if u['unit_id']==unit_id: return u
    return None

def image_path_of(obj):
    if not isinstance(obj,dict): return ''
    for k in ('file_path','path','url','src','video_path','image_path','thumbnail_path'):
        if obj.get(k): return str(obj.get(k))
    return ''

def asset_path(a):
    return image_path_of(a)

def is_video_path(p):
    return Path(str(p).split('?')[0]).suffix.lower() in VIDEO_EXTS

def media_type_of(obj, default='image'):
    mt=(obj.get('media_type') or obj.get('type') or '').lower() if isinstance(obj,dict) else ''
    if 'video' in mt: return 'video'
    p=image_path_of(obj)
    return 'video' if is_video_path(p) else default

def visible_candidates(shot):
    return [c for c in (shot.get('storyboard_candidates') or []) if not c.get('review_deleted')]

def project_stats(project, module_type):
    if module_type=='material':
        units=material_units(project)
        return {'unit_count':len(units),'asset_count':len(project.get('assets') or []),'scene_count':0,'shot_count':0}
    scs=project.get('scenes') or []
    shots=sum(len(s.get('shots') or []) for s in scs)
    cands=sum(len(visible_candidates(sh)) for s in scs for sh in (s.get('shots') or []))
    return {'scene_count':len(scs),'shot_count':shots,'candidate_count':cands,'asset_count':len(project.get('assets') or [])}

def resolve_file_path(wdir, raw):
    if not raw: return None
    wdir=Path(wdir)
    p=str(raw).replace('\\','/').strip()
    p=p.split('?',1)[0]
    p=p.lstrip('/')
    roots=[wdir/'extracted', wdir/'review_notes']
    for root in roots:
        cand=root/p
        if cand.exists() and cand.is_file(): return cand
    suffixes=[p]
    # 常见制作端导出：project.json 里是 /input 或 /output，但 zip 中可能多了 tools/image 前缀。
    for marker in ('input/','output/image/','output/video/','output/material/','temp_refs/','review_notes/'):
        if marker in p:
            suffixes.append(p[p.find(marker):])
    basename=Path(p).name
    for root in roots:
        if not root.exists(): continue
        for fp in root.rglob(basename):
            if not fp.is_file(): continue
            s=fp.as_posix().replace('\\','/')
            if any(s.endswith('/'+suf) or s.endswith(suf) for suf in suffixes):
                return fp
        # 最后兜底：只按 basename 找第一个。
        for fp in root.rglob(basename):
            if fp.is_file(): return fp
    return None

def reconcile_notes_shot_to_scene(project):
    """反向互通：制作端把审核备注写在子栏(shot)层，统一审核端按 SC/命名框(scene)层显示。
    导入时把子栏备注上提到所属 SC（仅当该 SC 本轮无备注时填充，取首个有备注的子栏）。幂等。"""
    if not isinstance(project, dict): return
    scenes=[]
    wss=project.get('material_workspaces') if isinstance(project.get('material_workspaces'),dict) else {}
    for key in ('character','scene','object'):
        for s in (wss.get(key) or {}).get('scenes') or []: scenes.append(s)
    for s in (project.get('scenes') or []): scenes.append(s)
    for sc in scenes:
        nbr=sc.setdefault('review_notes_by_round',{})
        rds=sc.setdefault('review',{}).setdefault('rounds',{})
        for sh in (sc.get('shots') or []):
            shnbr=sh.get('review_notes_by_round')
            if not isinstance(shnbr,dict): continue
            for r,shslot in shnbr.items():
                if not isinstance(shslot,dict): continue
                txt=shslot.get('text') or ''; imgs=shslot.get('images') or []
                if not txt and not imgs: continue
                cur=nbr.get(str(r)) or {}
                sc_has=bool(cur.get('images')) or bool(cur.get('text')) or bool((rds.get(str(r)) or {}).get('note'))
                if sc_has: continue
                slot=nbr.setdefault(str(r),{'text':'','images':[]})
                if imgs and not slot.get('images'): slot['images']=[dict(im) for im in imgs]
                if txt and not slot.get('text'): slot['text']=txt
                rd=rds.setdefault(str(r),{})
                if txt and not rd.get('note'): rd['note']=txt

def update_round(review, round_no, status, reviewer, note=''):
    ensure_rounds(review)
    rn=max(1,min(3,int(round_no or review.get('current_round') or 1)))
    if rn > int(review.get('max_round_available') or 1):
        review['max_round_available']=rn
    for i in range(1,rn+1):
        review['rounds'].setdefault(str(i), {'label':review_label(i),'status':'pending','locked':False,'submitted_at':'','submitted_by':'','note':''})
    review['current_round']=rn
    r=review['rounds'][str(rn)]
    r['status']=status if status in ('approved','rejected','pending') else 'pending'
    r['locked']=status in ('approved','rejected')
    r['submitted_at']=now_str()
    r['submitted_by']=reviewer or '审核员'
    r['note']=note or r.get('note','')
    if status=='rejected' and rn<3:
        nxt=rn+1
        review['rounds'].setdefault(str(nxt), {'label':review_label(nxt),'status':'pending','locked':False,'submitted_at':'','submitted_by':'','note':''})
        review['max_round_available']=max(int(review.get('max_round_available') or 1),nxt)
        review['current_round']=nxt
    return review

def make_review_result(project, manifest, usage, module_type):
    result={'schema_version':'2.0','module_type':module_type,'module_label':MODULE_LABELS.get(module_type,module_type),'exported_at':now_str(),'project_id':project.get('project_id'),'project_name':project.get('project_name'),'review_units':[]}
    if module_type=='material':
        ensure_material_review(project)
        units=project.get('material_review',{}).get('units',{})
        for u in material_units(project):
            result['review_units'].append({'unit_id':u['unit_id'],'unit_type':u['unit_type'],'category':u['category'],'title':u['title'],'asset_count':len(u.get('assets') or []),'review':units.get(u['unit_id'],{})})
    else:
        for sc in sorted(project.get('scenes') or [], key=scene_number):
            rv=ensure_scene_review(sc)
            result['review_units'].append({'unit_id':sc.get('scene_id') or scene_code(sc),'unit_type':'scene','title':scene_code(sc),'scene_code':scene_code(sc),'review':rv})
    return result

def compact_for_export(project):
    # 保持结果完整，不做破坏性压缩；只打标记。
    p=copy.deepcopy(project)
    p.setdefault('review_export_info',{})
    p['review_export_info'].update({'exported_from':'unified_review_client','exported_at':now_str()})
    return p

def generator_display(project, manifest=None):
    project = project or {}; manifest = manifest or {}
    names = project.get('generator_names') or manifest.get('generator_names')
    out = [str(x) for x in names if x] if isinstance(names, list) else []
    if not out:
        single = (project.get('generator_name') or manifest.get('generator_name')
                  or (project.get('export_owner') or {}).get('responsible_nickname')
                  or manifest.get('responsible_nickname') or '')
        if single: out = [single]
    return '、'.join(out)


def export_excel_bytes(project, usage, module_type, manifest=None):
    if openpyxl is None:
        raise RuntimeError('缺少 openpyxl，请先运行“安装”脚本安装依赖')
    wb=openpyxl.Workbook()
    ws=wb.active
    ws.title='审核表'
    green='00E5A0'; mid='1F2937'; white='FFFFFF'
    thin=Side(style='thin', color='334155')
    border=Border(left=thin,right=thin,top=thin,bottom=thin)
    header_fill=PatternFill('solid', fgColor=green)
    sub_fill=PatternFill('solid', fgColor=mid)
    pass_map={'approved':'通过','rejected':'不通过','pending':'待审核'}
    gen_disp=generator_display(project, manifest)
    def style_range(rng, fill=None, font=None, align='center'):
        for row in ws[rng]:
            for c in row:
                c.border=border
                c.alignment=Alignment(horizontal=align, vertical='center', wrap_text=True)
                if fill: c.fill=fill
                if font: c.font=font
    def setup_header(first_title='分镜序号'):
        ws.merge_cells('A1:B3'); ws.merge_cells('C1:C3'); ws.merge_cells('D1:I1')
        ws.merge_cells('D2:E2'); ws.merge_cells('F2:G2'); ws.merge_cells('H2:I2'); ws.merge_cells('J1:J3'); ws.merge_cells('K1:K3')
        ws['A1']=first_title; ws['C1']='生成时间'; ws['D1']='审核'; ws['D2']='一审'; ws['F2']='二审'; ws['H2']='三审'
        ws['D3']='审核时间'; ws['E3']='通过'; ws['F3']='审核时间'; ws['G3']='通过'; ws['H3']='审核时间'; ws['I3']='通过'; ws['J1']='费用'; ws['K1']='生成者'
        style_range('A1:K3', header_fill, Font(bold=True,color='00120C'), 'center')
        for r in (1,2,3): ws.row_dimensions[r].height=24
        for i,w in enumerate([16,12,22,22,12,22,12,22,12,14,18],1): ws.column_dimensions[get_column_letter(i)].width=w
        ws.freeze_panes='A4'
    setup_header('审核对象')
    row=4; total_cost=usage_total_usd(usage); scene_costs=usage_scene_usd_map(usage)
    if module_type=='material':
        ensure_material_review(project); reviews=project.get('material_review',{}).get('units',{})
        for u in material_units(project):
            rv=reviews.get(u['unit_id'],{}); rounds=rv.get('rounds') or {}
            ws.cell(row,1).value=u['category']; ws.cell(row,2).value=u['title']; ws.cell(row,3).value=first_time_from_items(u.get('assets') or [])
            for idx,col in enumerate((4,6,8),1):
                r=rounds.get(str(idx),{})
                ws.cell(row,col).value=r.get('submitted_at',''); ws.cell(row,col+1).value=pass_map.get(r.get('status','pending'),'待审核')
            ws.cell(row,10).value=total_cost if row==4 else ''
            ws.cell(row,11).value=gen_disp
            row+=1
    else:
        for sc in sorted(project.get('scenes') or [], key=scene_number):
            shots=sc.get('shots') or []
            if not shots: shots=[{'shot_code':'','tabs':[],'storyboard_candidates':[]}]
            start=row; rv=ensure_scene_review(sc); rounds=rv.get('rounds') or {}
            for sh in shots:
                ws.cell(row,1).value=scene_code(sc); ws.cell(row,2).value=sh.get('shot_code') or sh.get('shot_id') or ''; ws.cell(row,3).value=shot_generation_time(sh)
                for idx,col in enumerate((4,6,8),1):
                    r=rounds.get(str(idx),{})
                    ws.cell(row,col).value=r.get('submitted_at',''); ws.cell(row,col+1).value=pass_map.get(r.get('status','pending'),'待审核')
                ws.cell(row,11).value=gen_disp
                row+=1
            end=row-1
            if end>start:
                ws.merge_cells(start_row=start,start_column=1,end_row=end,end_column=1)
                ws.merge_cells(start_row=start,start_column=10,end_row=end,end_column=10)
                ws.merge_cells(start_row=start,start_column=11,end_row=end,end_column=11)
                ws.cell(start,11).value=gen_disp
            ws.cell(start,10).value=scene_costs.get(scene_code(sc), '')
    if row==4:
        ws.cell(4,1).value='暂无数据'; ws.merge_cells('A4:J4'); row=5
    for rr in range(4,row):
        for cc in range(1,12):
            c=ws.cell(rr,cc); c.border=border; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.row_dimensions[rr].height=28; ws.cell(rr,10).number_format='$0.000000'
    row += 1
    ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=9)
    ws.cell(row,1).value='总成本'; ws.cell(row,10).value=total_cost; ws.cell(row,10).number_format='$0.000000'
    style_range(f'A{row}:K{row}', sub_fill, Font(bold=True,color=white), 'center')
    bio=io.BytesIO(); wb.save(bio); return bio.getvalue()


def build_reviewed_package(wid):
    w,info,project,manifest,usage=load_workspace(wid)
    module_type=info.get('module_type') or detect_module(manifest,project,info.get('project_json_rel',''))
    extract=w/'extracted'
    project_json_rel=info.get('project_json_rel') or manifest.get('project_json') or 'project.json'
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as zf:
        written=set()
        if extract.exists():
            for fp in extract.rglob('*'):
                if not fp.is_file(): continue
                rel=fp.relative_to(extract).as_posix()
                if rel in {'review_result.json','review_state.json','manifest.json'}: continue
                if rel==project_json_rel:
                    continue
                zf.write(fp, rel)
                written.add(rel)
        # 审核端新增的备注图存放在 <工作区>/review_notes/（不在 extracted 下），
        # 必须一并打进包，否则下载方（制作端 / 其它审核端）拿不到图片，
        # 表现为「文字备注保留、图片损坏 / 丢失」。
        notes_dir=w/'review_notes'
        if notes_dir.exists():
            for fp in notes_dir.rglob('*'):
                if not fp.is_file(): continue
                arc='review_notes/'+fp.relative_to(notes_dir).as_posix()
                if arc in written: continue
                zf.writestr(arc, fp.read_bytes())
                written.add(arc)
        compact=compact_for_export(project)
        pinf=parent_info_from_workspace(info, project, manifest)
        compact['module_type']=module_type
        compact['module_label']=MODULE_LABELS.get(module_type,module_type)
        compact['module_ui_label']=MODULE_UI_LABELS.get(module_type,module_type)
        compact['parent_id']=pinf['parent_id']; compact['parent_name']=pinf['parent_name']
        compact['child_id']=pinf['child_id']; compact['child_name']=pinf['child_name']
        compact['project_scope']={'parent_id':pinf['parent_id'],'parent_name':pinf['parent_name'],'module_type':module_type,'module_label':MODULE_LABELS.get(module_type,module_type),'module_ui_label':MODULE_UI_LABELS.get(module_type,module_type),'child_id':pinf['child_id'],'child_name':pinf['child_name']}
        zf.writestr(project_json_rel, json.dumps(compact,ensure_ascii=False,indent=2).encode('utf-8'))
        manifest2=copy.deepcopy(manifest or {})
        manifest2['module_type']=module_type
        manifest2['module_label']=MODULE_LABELS.get(module_type,module_type)
        manifest2['module_ui_label']=MODULE_UI_LABELS.get(module_type,module_type)
        manifest2['parent_id']=pinf['parent_id']; manifest2['parent_name']=pinf['parent_name']
        manifest2['child_id']=pinf['child_id']; manifest2['child_name']=pinf['child_name']
        manifest2['project_scope']={'parent_id':pinf['parent_id'],'parent_name':pinf['parent_name'],'module_type':module_type,'module_label':MODULE_LABELS.get(module_type,module_type),'module_ui_label':MODULE_UI_LABELS.get(module_type,module_type),'child_id':pinf['child_id'],'child_name':pinf['child_name']}
        manifest2['reviewed_by_unified_client']=True
        manifest2['reviewed_at']=now_str()
        manifest2['project_json']=project_json_rel
        zf.writestr('manifest.json', json.dumps(manifest2,ensure_ascii=False,indent=2).encode('utf-8'))
        zf.writestr('review_result.json', json.dumps(make_review_result(project,manifest2,usage,module_type),ensure_ascii=False,indent=2).encode('utf-8'))
        zf.writestr('review_state.json', json.dumps({'workspace':wid,'module_type':module_type,'exported_at':now_str()},ensure_ascii=False,indent=2).encode('utf-8'))
    return out.getvalue()

def import_review_zip(data, filename=''):
    wid='review_'+uuid.uuid4().hex[:12]
    w=workspace_dir(wid); extracted=w/'extracted'; w.mkdir(parents=True,exist_ok=True)
    tmp=w/'upload.zip'; tmp.write_bytes(data)
    safe_extract_zip(tmp,extracted)
    manifest=load_manifest(extracted)
    pjson, rel=find_project_json(extracted,manifest)
    project=read_json(pjson,{}) or {}
    usage=read_json(extracted/'usage_summary.json',{}) or {}
    parent_data=read_parent_json_from_extract(extracted,manifest)
    module_type=detect_module(manifest,project,rel)
    scope=resolve_review_import_scope(manifest,project,parent_data,filename)
    child_id, child_name = choose_review_child_identity(scope['parent_id'], module_type, scope.get('source_child_id'), scope.get('source_child_name'))
    merge_review_parent(scope['parent_id'], scope['parent_name'])
    project['module_type']=module_type
    project['module_label']=MODULE_LABELS.get(module_type,module_type)
    project['module_ui_label']=MODULE_UI_LABELS.get(module_type,module_type)
    project['parent_id']=scope['parent_id']; project['parent_name']=scope['parent_name']
    project['child_id']=child_id; project['child_name']=child_name
    project['project_name']=child_name
    project['project_scope']={'parent_id':scope['parent_id'],'parent_name':scope['parent_name'],'module_type':module_type,'module_label':MODULE_LABELS.get(module_type,module_type),'module_ui_label':MODULE_UI_LABELS.get(module_type,module_type),'child_id':child_id,'child_name':child_name}
    project['review_import_info']={'imported_at':now_str(),'source_filename':filename,'source_parent_id':scope.get('source_parent_id'),'source_parent_name':scope.get('source_parent_name'),'target_parent_id':scope['parent_id'],'target_parent_name':scope['parent_name'],'parent_matched_by':scope.get('matched_by'),'created_new_parent':scope.get('created_new_parent'),'source_child_id':scope.get('source_child_id'),'source_child_name':scope.get('source_child_name'),'import_mode':'copy_no_overwrite'}
    if module_type=='material':
        ensure_material_review(project)
        project['__material_units']=material_units(project)
    else:
        for sc in project.get('scenes') or []: ensure_scene_review(sc)
    reconcile_notes_shot_to_scene(project)
    manifest['module_type']=module_type
    manifest['module_label']=MODULE_LABELS.get(module_type,module_type)
    manifest['module_ui_label']=MODULE_UI_LABELS.get(module_type,module_type)
    manifest['parent_id']=scope['parent_id']; manifest['parent_name']=scope['parent_name']
    manifest['child_id']=child_id; manifest['child_name']=child_name
    manifest['project_scope']={'parent_id':scope['parent_id'],'parent_name':scope['parent_name'],'module_type':module_type,'module_label':MODULE_LABELS.get(module_type,module_type),'module_ui_label':MODULE_UI_LABELS.get(module_type,module_type),'child_id':child_id,'child_name':child_name}
    if not manifest.get('project_json'): manifest['project_json']=rel
    imported_at=now_str()
    info={'workspace':wid,'source_filename':filename,'imported_at':imported_at,'project_json_rel':rel,'module_type':module_type,'parent_id':scope['parent_id'],'parent_name':scope['parent_name'],'source_parent_id':scope.get('source_parent_id'),'source_parent_name':scope.get('source_parent_name'),'child_id':child_id,'child_name':child_name,'source_child_id':scope.get('source_child_id'),'source_child_name':scope.get('source_child_name'),'parent_matched_by':scope.get('matched_by'),'created_new_parent':scope.get('created_new_parent')}
    write_json(w/'workspace.json',info); write_json(w/'manifest.json',manifest); write_json(w/'usage_summary.json',usage); write_json(w/'project.json',project)
    return {'ok':True,'workspace':wid,'module_type':module_type,'module_label':MODULE_LABELS.get(module_type,module_type),'module_ui_label':MODULE_UI_LABELS.get(module_type,module_type),'parent_id':scope['parent_id'],'parent_name':scope['parent_name'],'child_id':child_id,'child_name':child_name,'project_name':child_name}


INDEX_HTML = r"""<!doctype html><html><head><meta charset="utf-8"><title>帝蓝工作流-统一审核端</title><style>
*{box-sizing:border-box}html,body{height:100%;margin:0;background:#080b0f;color:#e8eaed;font-family:'Microsoft YaHei',Arial,sans-serif;overflow:hidden}button,input,select,textarea{font-family:inherit}.app{height:100vh;display:flex;flex-direction:column}.top{height:58px;display:flex;align-items:center;gap:10px;padding:0 14px;background:#080b0f;border-bottom:1px solid rgba(255,255,255,.08);flex-shrink:0}.brand{font-weight:900;font-size:17px;display:flex;align-items:center;gap:8px;white-space:nowrap}.brand i{width:27px;height:27px;background:#00e5a0;color:#00120c;border-radius:8px;display:inline-flex;align-items:center;justify-content:center;font-style:normal}.spacer{flex:1}.btn{border:0;border-radius:10px;background:#00e5a0;color:#00120c;font-weight:900;padding:8px 12px;cursor:pointer}.btn:hover{filter:brightness(1.06)}.btn.ghost{background:#171d24;color:#e8eaed;border:1px solid rgba(255,255,255,.13)}.btn.danger{background:rgba(255,92,92,.1);color:#ff7474;border:1px solid rgba(255,92,92,.35)}.btn.small{padding:5px 8px;font-size:12px}.nav{display:flex;gap:9px;align-items:center}.top-nav-btn{height:34px;border-radius:10px;padding:7px 13px}.nav button.active,#navManager.active{border-color:rgba(0,229,160,.55);color:#00e5a0;background:rgba(0,229,160,.1);box-shadow:0 0 0 1px rgba(0,229,160,.08) inset}.nav button{white-space:nowrap}.pill{font-size:12px;color:#9aa3ad;background:#11161d;border:1px solid rgba(255,255,255,.1);border-radius:999px;padding:6px 10px;white-space:nowrap}.page{flex:1;min-height:0;overflow:auto;padding:16px}.hidden{display:none!important}.manager-layout{display:grid;grid-template-columns:300px 1fr;gap:14px;height:100%}.manager-side,.manager-main{background:#0f141a;border:1px solid rgba(255,255,255,.09);border-radius:18px;padding:14px;min-height:0;overflow:auto}.page-title{font-size:22px;font-weight:900;margin-bottom:6px}.desc{color:#8a9098;font-size:13px}.project-grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(300px,1fr));gap:12px;margin-top:14px}.project-card{background:#111720;border:1px solid rgba(255,255,255,.1);border-radius:16px;padding:14px;cursor:pointer}.project-card:hover{border-color:rgba(0,229,160,.45);transform:translateY(-1px)}.project-card h3{margin:0 0 8px;font-size:17px}.project-meta{font-size:12px;color:#8a9098;line-height:1.7}.module-filter{display:flex;flex-direction:column;gap:8px;margin-top:14px}.module-filter button{text-align:left}.subbar{height:44px;display:flex;align-items:center;gap:8px;padding:0 16px;background:#0b0f14;border-bottom:1px solid rgba(255,255,255,.08)}.cat-tabs{display:flex;gap:8px}.cat-tabs button{background:#151b23;color:#d5d9df;border:1px solid rgba(255,255,255,.12);border-radius:10px;padding:8px 16px;font-weight:900;cursor:pointer}.cat-tabs button.active{color:#00120c;background:#00e5a0;border-color:#00e5a0}.scene,.unit{background:#0f141a;border:1px solid rgba(255,255,255,.09);border-radius:18px;margin-bottom:15px;overflow:hidden}.scene-head,.unit-head{display:flex;align-items:center;gap:8px;min-height:56px;padding:10px 12px;background:#131922;border-bottom:1px solid rgba(255,255,255,.07);flex-wrap:wrap}.sc-code,.unit-title{font-weight:900;font-size:16px;min-width:100px}.review-tools{margin-left:auto;display:flex;gap:8px;align-items:center;flex-wrap:wrap}.status-label{font-weight:900;font-size:13px;padding:5px 9px;border-radius:999px}.status-pending{background:rgba(245,166,35,.12);color:#f7c568;border:1px solid rgba(245,166,35,.36)}.status-approved{background:rgba(0,229,160,.12);color:#00e5a0;border:1px solid rgba(0,229,160,.36)}.status-rejected{background:rgba(255,92,92,.12);color:#ff7474;border:1px solid rgba(255,92,92,.36)}select,input,textarea{background:#171d24;border:1px solid rgba(255,255,255,.13);color:#e8eaed;border-radius:8px;padding:6px;outline:none}.strip{display:flex;gap:12px;align-items:flex-start;padding:13px;overflow-x:auto}.card{width:190px;flex:0 0 190px}.card-title{font-size:13px;color:#d5d9df;font-weight:800;text-align:center;margin-bottom:6px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}.media{height:112px;background:#050607;border:1px solid rgba(255,255,255,.1);border-radius:13px;display:flex;align-items:center;justify-content:center;overflow:hidden;position:relative}.media img,.media video{max-width:100%;max-height:100%;object-fit:contain}.empty{font-size:12px;color:#596271;text-align:center;line-height:1.5}.foot{height:32px;display:flex;align-items:center;justify-content:center;gap:8px;color:#9aa3ad}.arrow{background:#171d24;color:#d5d9df;border:1px solid rgba(255,255,255,.13);border-radius:8px;min-width:30px;height:26px;cursor:pointer}.note-row{display:flex;gap:8px;align-items:flex-start;padding:0 13px 13px}.note-row textarea{flex:1;min-height:52px;resize:vertical}.story-page{height:calc(100vh - 58px);background:#000;position:relative;overflow:hidden}.story-stage{position:absolute;inset:0;display:flex;align-items:center;justify-content:center;background:#000}.story-stage img,.story-stage video{width:100%;height:100%;object-fit:contain;background:#000}.story-black{width:100%;height:100%;background:#000;display:flex;align-items:center;justify-content:center;color:#777;font-size:22px}.player-bar{position:absolute;left:18px;right:18px;bottom:18px;background:rgba(12,15,19,.78);backdrop-filter:blur(12px);border:1px solid rgba(255,255,255,.12);border-radius:16px;padding:11px;display:flex;align-items:center;gap:10px}.player-info{color:#d5d9df;font-size:13px;min-width:260px;text-align:right}.progress{height:12px;background:#333;border-radius:999px;overflow:hidden;flex:1;min-width:260px;cursor:pointer;touch-action:none}.progress span{display:block;height:100%;background:#00e5a0;width:0}.readonly{opacity:.86}.prod-workspace{display:grid;grid-template-columns:1fr 330px;gap:14px;height:100%}.prod-main,.prod-side{min-height:0;overflow:auto}.prod-scene{background:#0f141a;border:1px solid rgba(255,255,255,.09);border-radius:18px;margin-bottom:15px;overflow:hidden}.prod-scene-head{display:flex;align-items:center;gap:8px;min-height:48px;padding:10px 12px;background:#131922;border-bottom:1px solid rgba(255,255,255,.07)}.prod-shot{padding:12px;border-top:1px solid rgba(255,255,255,.05)}.prod-shot-head{display:flex;align-items:center;gap:8px;margin-bottom:8px}.prod-tabs{display:flex;gap:8px;flex-wrap:wrap;margin-bottom:8px}.prod-tab{padding:6px 10px;border-radius:999px;background:#171d24;border:1px solid rgba(255,255,255,.1);font-size:12px}.flow-grid,.asset-grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(150px,1fr));gap:10px}.asset-card{background:#10161d;border:1px solid rgba(255,255,255,.1);border-radius:12px;padding:8px}.asset-card img,.asset-card video{width:100%;height:110px;object-fit:contain;background:#000;border-radius:8px}.meta{font-size:12px;color:#8a9098;margin-top:4px}.msg{border-left:3px solid rgba(0,229,160,.5);padding:8px 10px;background:#0b0f14;border-radius:8px;margin:6px 0;white-space:pre-wrap}.model-chip{display:inline-block;margin:2px 4px 2px 0;padding:3px 7px;border-radius:999px;background:rgba(0,229,160,.08);border:1px solid rgba(0,229,160,.25);color:#00e5a0;font-size:12px}.modal{position:fixed;inset:0;background:rgba(0,0,0,.72);display:none;align-items:center;justify-content:center;z-index:99}.modal.show{display:flex}.modal img,.modal video{max-width:92vw;max-height:90vh;object-fit:contain;background:#000;border-radius:12px}.sc-content{display:flex;gap:12px;align-items:stretch;padding:13px}.sc-content .strip{flex:1;min-width:0;padding:0}.note-side{width:300px;flex-shrink:0;display:flex;flex-direction:column;gap:8px;border-left:1px solid rgba(255,255,255,.08);padding-left:12px}.note-head{font-weight:800;font-size:13px;color:#cfd4da}.note-imgbox{border:1px dashed rgba(0,229,160,.35);border-radius:12px;padding:8px;min-height:96px;background:#0b0e12}.note-imgbox.drag{border-color:#00e5a0;background:rgba(0,229,160,.06)}.note-empty{font-size:12px;color:#596271;line-height:1.6}.note-text{width:100%;min-height:90px;resize:vertical;background:#0b0e12;color:#e8eaed;border:1px solid rgba(255,255,255,.12);border-radius:10px;padding:8px;outline:none}@media(max-width:1100px){.sc-content{flex-direction:column}.note-side{width:auto;border-left:0;border-top:1px solid rgba(255,255,255,.08);padding-left:0;padding-top:10px}}.note-imgs{display:flex;gap:6px;flex-wrap:wrap;align-items:center;margin-top:6px}.note-thumb{position:relative;width:56px;height:56px;border-radius:8px;overflow:hidden;border:1px solid rgba(255,255,255,.12)}.note-thumb img{width:100%;height:100%;object-fit:cover;cursor:pointer}.note-del{position:absolute;top:1px;right:1px;width:16px;height:16px;border:0;border-radius:50%;background:rgba(0,0,0,.6);color:#ff8080;font-size:11px;line-height:1;cursor:pointer;padding:0}
.cand-del{position:absolute;top:6px;right:6px;width:24px;height:24px;border:0;border-radius:50%;background:rgba(0,0,0,.62);color:#ff6b6b;font-size:16px;line-height:1;cursor:pointer;padding:0;display:none;z-index:5;box-shadow:0 1px 4px rgba(0,0,0,.45)}
.media:hover .cand-del{display:block}
.cand-del:hover{background:rgba(214,48,49,.92);color:#fff}
.pv-del{position:fixed;top:24px;right:28px;width:42px;height:42px;border:0;border-radius:50%;background:rgba(0,0,0,.62);color:#ff6b6b;font-size:24px;line-height:1;cursor:pointer;display:none;z-index:30;box-shadow:0 2px 8px rgba(0,0,0,.5)}
#preview.has-del:hover .pv-del{display:block}
.pv-del:hover{background:rgba(214,48,49,.95);color:#fff}.time-inp{width:46px;background:#0b0e12;border:1px solid rgba(255,255,255,.18);color:#e8eaed;border-radius:6px;padding:3px 4px;text-align:center;font-size:12px;margin:0 1px}.sc-time-edit{display:inline-flex;align-items:center;gap:2px;color:#8a9098;font-size:12px;flex-wrap:wrap}#reviewDirtyTip{display:none;color:#f5a623;font-size:12px;margin-left:6px}.toast{position:fixed;left:50%;bottom:22px;transform:translateX(-50%);background:#121a22;border:1px solid rgba(255,255,255,.16);border-radius:12px;padding:10px 14px;color:#fff;display:none;z-index:4000}.toast.show{display:block}.locked-tag{font-size:12px;color:#9aa3ad}

.manager-layout{display:grid;grid-template-columns:320px 1fr;gap:0;height:min(760px,calc(100vh - 110px));max-width:1120px;margin:14px auto;background:#0b1016;border:1px solid rgba(255,255,255,.12);border-radius:20px;overflow:hidden}.manager-side,.manager-main{border:0;border-radius:0;background:transparent}.manager-side{border-right:1px solid rgba(255,255,255,.09);padding:18px 20px}.manager-main{padding:0}.parent-create{display:flex;gap:8px;margin-bottom:10px}.parent-create input{flex:1;min-width:0}.parent-card{background:#0f161e;border:1px solid rgba(255,255,255,.11);border-radius:14px;padding:12px;margin-top:12px;cursor:pointer}.parent-card.active{border-color:rgba(0,229,160,.65);background:rgba(0,229,160,.08)}.parent-card b{display:block;margin-bottom:5px}.parent-card .actions{display:flex;gap:8px;margin-top:8px}.manager-module-row{height:74px;display:flex;align-items:center;gap:16px;padding:0 22px;border-bottom:1px solid rgba(255,255,255,.08)}.module-pill-btn{min-width:96px;height:44px;border-radius:999px;font-size:20px}.module-pill-btn.ghost{opacity:.72}.manager-import{margin-left:auto}.child-toolbar{display:flex;align-items:center;gap:10px;padding:18px 22px 8px}.child-list{padding:0 22px 22px}.child-card{display:flex;align-items:center;gap:12px;background:#111820;border:1px solid rgba(255,255,255,.1);border-radius:14px;padding:14px;margin-bottom:10px}.child-card .child-title{font-weight:900;font-size:16px}.child-card .child-main{flex:1;min-width:0}.child-card .child-actions{display:flex;gap:8px;align-items:center}.count-line{color:#9aa3ad;font-size:13px;line-height:1.65}.empty-center{padding:24px;color:#8a9098}.import-empty{height:100%;display:flex;align-items:center;justify-content:center}.import-card{background:#0f141a;border:1px solid rgba(255,255,255,.1);border-radius:18px;padding:24px;text-align:center;max-width:420px}


/* V6 制作流程只读镜像：复刻制作端主体工作台，不包含制作端顶部后端操作栏 */
#flowPage{padding:0;overflow:hidden}.flow-readonly{height:100%;display:flex;min-height:0;background:#0a0c0f}.flow-readonly *{box-sizing:border-box}.flow-readonly .workspace-main{flex:1;min-width:0;overflow:auto;padding:16px;background:#0a0c0f}.flow-readonly .side-panel{width:310px;background:#0d1014;border-left:1px solid rgba(255,255,255,.08);padding:12px;overflow:auto;flex-shrink:0}.flow-readonly .page-title{font-size:20px;font-weight:900;margin:0 0 4px}.flow-readonly .desc{font-size:13px;color:#8a9098;margin-bottom:14px}.flow-readonly .scene{border:1px solid rgba(255,255,255,.09);border-radius:18px;background:#0d1014;margin-bottom:16px;overflow:hidden}.flow-readonly .scene-head{min-height:48px;display:flex;align-items:center;gap:8px;padding:7px 12px;background:#11161d;border-bottom:1px solid rgba(255,255,255,.07);flex-wrap:wrap}.flow-readonly .scene-title{font-weight:900;border:1px solid transparent;border-radius:8px;padding:4px 6px}.flow-readonly .scene-time{display:flex;align-items:center;gap:4px;color:#8a9098;font-size:12px}.flow-readonly .scene-actions{margin-left:auto;display:flex;gap:7px}.flow-readonly .shot-list{padding:12px;display:flex;flex-direction:column;gap:14px}.flow-readonly .shot{background:#11161d;border:1px solid rgba(255,255,255,.09);border-radius:18px;overflow:hidden}.flow-readonly .shot-head{min-height:46px;display:flex;align-items:center;gap:8px;padding:7px 10px;border-bottom:1px solid rgba(255,255,255,.07);background:#121820;flex-wrap:wrap}.flow-readonly .handle{border:1px solid rgba(255,255,255,.1);border-radius:8px;padding:4px 7px;color:#596271}.flow-readonly .shot-title{font-weight:900}.flow-readonly .shot-duration,.flow-readonly .status{display:flex;align-items:center;gap:5px;color:#8a9098;font-size:12px}.flow-readonly .dot{width:8px;height:8px;border-radius:50%;display:inline-block}.flow-readonly .dot.red{background:#ff5c5c}.flow-readonly .dot.yellow{background:#f5a623}.flow-readonly .dot.green{background:#00e5a0}.flow-readonly .shot-tabs{display:flex;align-items:center;gap:5px;padding:8px 10px 0;background:#0f141b;overflow-x:auto;border-bottom:1px solid rgba(255,255,255,.06)}.flow-readonly .work-tab{height:34px;min-width:94px;max-width:190px;display:flex;align-items:center;gap:7px;background:#171d24;color:#8a9098;border:1px solid rgba(255,255,255,.1);border-bottom-color:transparent;border-radius:11px 11px 0 0;padding:0 8px;cursor:pointer;white-space:nowrap}.flow-readonly .work-tab.active{background:#0b0e12;color:#00e5a0;border-color:rgba(0,229,160,.32);border-bottom-color:#0b0e12}.flow-readonly .work-tab span{overflow:hidden;text-overflow:ellipsis}.flow-readonly .tab-menu{display:flex;gap:5px;margin-left:4px}.flow-readonly .shot-body{display:grid;grid-template-columns:minmax(620px,1fr) 390px;gap:12px;padding:12px}.flow-readonly .tabbed-work{min-width:0}.flow-readonly .tab-content{display:grid;grid-template-columns:235px minmax(360px,1fr);gap:12px}.flow-readonly .shot-left,.flow-readonly .shot-center,.flow-readonly .story{min-width:0}.flow-readonly .settings{display:flex;flex-direction:column;gap:8px;margin-bottom:10px}.flow-readonly .settings label{font-size:12px;color:#8a9098;display:flex;justify-content:space-between;align-items:center;gap:8px}.flow-readonly .ro-select,.flow-readonly .ro-input{min-width:96px;max-width:132px;background:#171d24;border:1px solid rgba(255,255,255,.12);color:#e8eaed;border-radius:8px;padding:6px;text-align:center;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}.flow-readonly .refs{border:1px solid rgba(255,255,255,.09);background:#0b0e12;border-radius:14px;padding:9px;min-height:150px}.flow-readonly .refs-head{display:flex;align-items:center;gap:6px;margin-bottom:8px}.flow-readonly .refs-title{font-size:12px;color:#8a9098;flex:1}.flow-readonly .refs-list{display:flex;flex-direction:column;gap:7px}.flow-readonly .ref-chip{display:flex;align-items:center;gap:6px;border:1px solid rgba(0,229,160,.2);background:rgba(0,229,160,.07);border-radius:10px;padding:5px}.flow-readonly .ref-chip img,.flow-readonly .ref-chip video,.flow-readonly .ref-chip .media-icon{width:38px;height:38px;object-fit:contain;background:#000;border-radius:7px;display:flex;align-items:center;justify-content:center;color:#8fdcff}.flow-readonly .ref-chip span{font-size:12px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}.flow-readonly .video-info{border:1px solid rgba(255,255,255,.09);background:#111820;border-radius:14px;padding:12px;color:#b9c0c9;line-height:1.6;margin-bottom:10px}.flow-readonly .video-info.success{border-color:rgba(0,229,160,.25);background:rgba(0,229,160,.06)}.flow-readonly .video-info.error{border-color:rgba(255,92,92,.32);background:rgba(255,92,92,.07)}.flow-readonly .composer{display:grid;grid-template-columns:1fr auto;gap:8px;align-items:start}.flow-readonly .prompt-area{min-width:0}.flow-readonly .prompt{width:100%;min-height:170px;background:#0b0e12;border:1px solid rgba(255,255,255,.12);color:#e8eaed;border-radius:12px;padding:10px;resize:vertical;line-height:1.6;white-space:pre-wrap;overflow:auto}.flow-readonly .composer-actions{display:flex;align-items:flex-start}.flow-readonly .disabled-action{opacity:.48;cursor:not-allowed;filter:grayscale(.2)}.flow-readonly .image-grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(170px,1fr));gap:8px;margin-top:10px}.flow-readonly .gen-img{position:relative;background:#050607;border:1px solid rgba(255,255,255,.09);border-radius:12px;overflow:hidden}.flow-readonly .gen-img img,.flow-readonly .gen-img video{width:100%;height:130px;object-fit:contain;background:#000;display:block}.flow-readonly .img-caption{font-size:12px;color:#8a9098;padding:6px 8px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}.flow-readonly .msg{border:0;padding:0;background:transparent;border-radius:0;margin:8px 0;white-space:normal}.flow-readonly .msg-role{font-size:12px;color:#8a9098;margin-bottom:4px}.flow-readonly .bubble{border-left:3px solid rgba(0,229,160,.5);padding:8px 10px;background:#0b0f14;border-radius:8px;white-space:pre-wrap;line-height:1.55}.flow-readonly .story{background:#0b0e12;border:1px solid rgba(255,255,255,.09);border-radius:16px;overflow:hidden;display:flex;flex-direction:column;min-height:330px}.flow-readonly .story-head{min-height:38px;display:flex;align-items:center;gap:8px;padding:8px 10px;border-bottom:1px solid rgba(255,255,255,.07);background:#10161d}.flow-readonly .story-switch{display:flex;gap:7px}.flow-readonly .story-preview,.flow-readonly .note-preview{height:230px;background:#000;display:flex;align-items:center;justify-content:center;overflow:hidden}.flow-readonly .story-preview img,.flow-readonly .story-preview video,.flow-readonly .note-preview img,.flow-readonly .note-preview video{max-width:100%;max-height:100%;object-fit:contain}.flow-readonly .story-empty{color:#596271;text-align:center;line-height:1.7}.flow-readonly .story-foot,.flow-readonly .note-tools{height:42px;display:flex;align-items:center;justify-content:center;gap:10px;border-top:1px solid rgba(255,255,255,.07)}.flow-readonly .note-box{width:100%;min-height:86px;background:#0b0e12;border:0;border-top:1px solid rgba(255,255,255,.07);border-radius:0;color:#e8eaed;resize:vertical}.flow-readonly .side-tabs{display:flex;gap:8px;margin-bottom:10px}.flow-readonly .side-tabs button,.flow-readonly .tab-btn{flex:1;background:#171d24;color:#8a9098;border:1px solid rgba(255,255,255,.12);border-radius:10px;padding:8px;font-weight:800;cursor:pointer}.flow-readonly .side-tabs button.active,.flow-readonly .tab-btn.active{color:#00e5a0;background:rgba(0,229,160,.08);border-color:rgba(0,229,160,.32)}.flow-readonly .asset-tabs{display:flex;gap:6px;margin:8px 0;flex-wrap:wrap}.flow-readonly .asset-tabs .tab-btn{flex:0 0 auto;padding:8px 10px}.flow-readonly .asset-grid{display:grid;grid-template-columns:repeat(2,1fr);gap:8px}.flow-readonly .asset-card,.flow-readonly .library-card{background:#11161d;border:1px solid rgba(255,255,255,.09);border-radius:13px;overflow:hidden;position:relative}.flow-readonly .asset-card img,.flow-readonly .library-card img,.flow-readonly .asset-card video,.flow-readonly .library-card video,.flow-readonly .asset-card .media-icon,.flow-readonly .library-card .media-icon{width:100%;height:92px;object-fit:contain;background:#050607;display:block}.flow-readonly .asset-card .media-icon,.flow-readonly .library-card .media-icon{font-size:34px;color:#8fdcff;line-height:92px;text-align:center}.flow-readonly .asset-card .row,.flow-readonly .library-row{display:flex;gap:6px;align-items:center;padding:6px}.flow-readonly .asset-card .name{flex:1;min-width:0;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;font-size:12px}.flow-readonly .mini{border:0;background:transparent;color:#8fdcff;cursor:pointer;padding:2px}.flow-readonly .mini.del{color:#ff7777}.flow-readonly .group-tools{display:flex;gap:8px;align-items:center;margin:8px 0}.flow-readonly .asset-groups{display:flex;flex-direction:column;gap:10px}.flow-readonly .asset-group{background:#0f141b;border:1px solid rgba(255,255,255,.09);border-radius:14px;overflow:hidden}.flow-readonly .asset-group-head{display:flex;justify-content:space-between;gap:8px;align-items:center;padding:8px;border-bottom:1px solid rgba(255,255,255,.07);background:#111820}.flow-readonly .asset-group-top{display:flex;gap:8px;align-items:center;min-width:0}.flow-readonly .asset-group-title{overflow:hidden;text-overflow:ellipsis;white-space:nowrap}.flow-readonly .asset-group-controls{display:flex;gap:7px;align-items:center;flex-wrap:wrap}.flow-readonly .group-grid{padding:8px}.flow-readonly .drop-hint{grid-column:1/-1;border:1px dashed rgba(0,229,160,.25);border-radius:12px;padding:9px;color:#8a9098;font-size:12px;text-align:center}.flow-readonly .board-grid{display:flex;flex-direction:column;gap:10px}.flow-readonly .board-card{background:#11161d;border:1px solid rgba(255,255,255,.09);border-radius:13px;padding:9px}.flow-readonly .board-card h4{margin:0 0 8px;font-size:13px}.flow-readonly .board-img{height:130px;background:#050607;border-radius:10px;display:flex;align-items:center;justify-content:center;overflow:hidden}.flow-readonly .board-img img,.flow-readonly .board-img video{max-width:100%;max-height:100%;object-fit:contain}.flow-readonly .board-foot{display:flex;align-items:center;justify-content:center;gap:8px;margin-top:7px}@media(max-width:1200px){.flow-readonly .shot-body{grid-template-columns:1fr}.flow-readonly .story{min-height:360px}.flow-readonly .tab-content{grid-template-columns:1fr}.flow-readonly .side-panel{display:none}}
.modal audio{max-width:92vw;width:70vw;background:#111;border-radius:12px}

</style></head><body><div class="app"><header class="top"><div class="brand"><i>审</i><span id="brandTitle">帝蓝工作流-统一审核端</span></div><button id="navManager" class="btn ghost top-nav-btn" onclick="showManager()">项目管理</button><div class="nav" id="pageNav"><button class="btn ghost top-nav-btn" id="navReview" onclick="showPage('review')">审核界面</button><button class="btn ghost top-nav-btn" id="navStory" onclick="showPage('story')">动态分镜</button><button class="btn ghost top-nav-btn" id="navFlow" onclick="showPage('flow')">制作流程</button></div><span class="pill" id="projectInfo">项目管理</span><div class="spacer"></div><input type="file" id="importFile" class="hidden" accept=".zip,application/zip" onchange="importPackage(this.files&&this.files[0]);this.value=''"/><button class="btn" onclick="dcUploadProject()">上传项目</button><button class="btn ghost" onclick="dcOpenDownload()">下载项目</button><button class="btn ghost" onclick="$('importFile').click()">本地导入</button><button class="btn ghost" onclick="dcConfig()">数据中心地址</button><button class="btn ghost" onclick="exportExcel()">表单导出</button><button class="btn" onclick="exportPackage()">审核结果导出</button></header><div id="subbar" class="subbar hidden"></div><main id="managerPage" class="page"></main><main id="reviewPage" class="page hidden"></main><main id="storyPage" class="story-page hidden"><div id="storyStage" class="story-stage"></div><div class="player-bar"><button class="btn small" id="playBtn" onclick="togglePlay()">播放</button><button class="btn ghost small" onclick="restartStory()">重播</button><div class="progress" id="storyProgress" onpointerdown="beginSeek(event)"><span id="progressFill"></span></div><div class="player-info" id="playerInfo">00:00 / 00:00</div></div></main><main id="flowPage" class="page hidden"></main></div><div class="modal" id="preview" onclick="this.classList.remove('show');stopPreview()"><img id="previewImg" class="hidden"><video id="previewVideo" class="hidden" controls></video><audio id="previewAudio" class="hidden" controls></audio><button id="pvDel" class="pv-del" title="删除该图" onclick="event.stopPropagation();delPreviewCurrent()">×</button></div><div class="toast" id="toast"></div><script>
let screen='manager',projects=[],parents=[],currentParent=null,workspace=null,project=null,manifest={},usageSummary={},moduleType=null,currentModule='material',currentPage='review',materialCat='人物',flowSideMode='assets',flowSideCategory='人物';
let storyList=[],storyIdx=0,storyPlaying=false,storyTimer=null,storyCurrentMs=0,storyTotalMs=0,storyStartedAt=0,storyFrameOffsetMs=0,candidateIndex={},flowCandidateIndex={},flowNoteIndex={},flowTabState={},flowGroupOpen={},seeking=false,wasPlayingBeforeSeek=false;
const $=id=>document.getElementById(id);function esc(s){return String(s??'').replace(/[&<>"']/g,m=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[m]))}function toast(t){let el=$('toast');el.textContent=t;el.classList.add('show');setTimeout(()=>el.classList.remove('show'),2400)}
function api(url,opts={}){opts.headers=opts.headers||{};if(opts.body&&!(opts.body instanceof FormData)){opts.headers['Content-Type']='application/json';opts.body=JSON.stringify(opts.body)}return fetch(url,opts).then(async r=>{if(!r.ok)throw new Error(await r.text()||r.status);return r.headers.get('content-type')?.includes('json')?r.json():r})}
function labelOf(m){return {material:'美术模块',image:'分镜模块',video:'视频模块'}[m]||m}function shortLabel(m){return {material:'美术',image:'分镜',video:'视频'}[m]||m}
function reviewClientTitle(){if(screen!=='workspace'||!project)return '帝蓝工作流-统一审核端';let name={material:'美术',image:'分镜',video:'视频'}[moduleType]||shortLabel(moduleType);return `帝蓝工作流-${name}审核端`}
function updateTop(){let title=reviewClientTitle();$('brandTitle').textContent=title;document.title=title;$('navManager').classList.toggle('active',screen==='manager');$('pageNav').classList.toggle('hidden',screen==='manager');$('navReview').classList.toggle('active',currentPage==='review');$('navStory').classList.toggle('active',currentPage==='story');$('navFlow').classList.toggle('active',currentPage==='flow');$('navStory').classList.toggle('hidden',currentModule==='material');let pinf=currentParentItem();$('projectInfo').textContent=screen==='manager'?'项目管理':(project?(labelOf(moduleType)+' · '+(pinf?.parent_name?escText(pinf.parent_name)+' / ':'')+(project.child_name||project.project_name||project.project_id||'未命名')):'未导入')}
function escText(s){return String(s??'')}
function switchModule(m){if(screen==='workspace'&&project&&moduleType&&m!==moduleType){alert('当前已打开 '+labelOf(moduleType)+' 工程。如需审核其他模块，请回到项目管理或导入对应工程包。');return}currentModule=m;if(currentPage==='story'&&m==='material')currentPage='review';render()}
function showManager(){screen='manager';workspace=null;project=null;moduleType=null;currentPage='review';loadProjects()}function showPage(p){if(p==='story'&&currentModule==='material')return;currentPage=p;screen='workspace';if(p==='story')buildStoryList();render()}
async function loadProjects(){try{let r=await api('/api/projects');projects=r.projects||[];parents=r.parents||[];if(!currentParent&&parents.length)currentParent=parents[0].parent_id;if(currentParent&&!parents.some(p=>p.parent_id===currentParent))currentParent=parents[0]?.parent_id||null}catch(e){projects=[];parents=[];toast('读取项目列表失败：'+e.message)}render()}
function currentParentItem(){return parents.find(p=>p.parent_id===currentParent)||null}
function parentCountLine(p){let c=p.module_counts||{};return `美术 ${c.material||0} · 分镜 ${c.image||0} · 视频 ${c.video||0}`}
async function createParent(){let input=$('newParentName');let name=(input?.value||'').trim();if(!name){toast('请输入父项目名');return}try{let r=await api('/api/create_parent',{method:'POST',body:{parent_name:name}});currentParent=r.parent?.parent_id||currentParent;if(input)input.value='';await loadProjects()}catch(e){toast('新建父项目失败：'+e.message)}}
async function renameParent(pid){let p=parents.find(x=>x.parent_id===pid);let name=prompt('新的父项目名',p?.parent_name||pid);if(name===null)return;name=name.trim();if(!name)return;try{await api('/api/rename_parent',{method:'POST',body:{parent_id:pid,parent_name:name}});await loadProjects()}catch(e){toast('改名失败：'+e.message)}}
async function renameChild(wid){let item=projects.find(x=>x.workspace===wid);let cur=item?(item.child_name||item.project_name||''):'';let name=prompt('新的子项目名',cur);if(name===null)return;name=name.trim();if(!name)return;try{await api('/api/rename_child',{method:'POST',body:{workspace:wid,child_name:name}});await loadProjects();toast('已改名：'+name)}catch(e){toast('改名失败：'+e.message)}}
async function deleteParent(pid){let p=parents.find(x=>x.parent_id===pid);if(!confirm(`确定删除父项目【${p?.parent_name||pid}】及其下面所有本地审核工程？`))return;try{await api('/api/delete_parent',{method:'POST',body:{parent_id:pid}});if(currentParent===pid)currentParent=null;await loadProjects()}catch(e){toast('删除失败：'+e.message)}}
async function importPackage(file){if(!file)return;let fd=new FormData();fd.append('file',file);$('projectInfo').textContent='正在导入...';try{let data=await fetch('/api/import',{method:'POST',body:fd}).then(async r=>{if(!r.ok)throw new Error(await r.text());return r.json()});currentParent=data.parent_id||currentParent;currentModule=data.module_type||currentModule;if(currentModule==='material'&&currentPage==='story')currentPage='review';workspace=data.workspace;screen='workspace';currentPage='review';await loadWorkspace();await loadProjects();toast(`已导入到【${data.parent_name||''} / ${shortLabel(data.module_type)} / ${data.child_name||data.project_name||''}】`)}catch(e){toast('导入失败：'+e.message);render()}}
async function openWorkspace(wid){workspace=wid;screen='workspace';await loadWorkspace()}async function deleteWorkspace(wid){if(!confirm('确定删除这个本地审核工程？原始导出包不受影响。'))return;await api('/api/delete_workspace',{method:'POST',body:{workspace:wid}});await loadProjects()}
async function loadWorkspace(){let r=await api('/api/workspace?workspace='+encodeURIComponent(workspace));project=r.project;manifest=r.manifest||{};usageSummary=r.usage_summary||{};moduleType=r.module_type;currentModule=moduleType;if(r.info&&r.info.parent_id)currentParent=r.info.parent_id;if(currentModule==='material'&&currentPage==='story')currentPage='review';render()}
function syncNoteTextsBeforeRerender(){try{document.querySelectorAll('textarea.note-text[data-unit-type]').forEach(ta=>{writeNoteText(ta.dataset.unitType,ta.dataset.unitId,ta.dataset.round,ta.value)})}catch(e){}}
function render(){syncNoteTextsBeforeRerender();updateTop();$('managerPage').classList.toggle('hidden',screen!=='manager');$('reviewPage').classList.toggle('hidden',screen!=='workspace'||currentPage!=='review');$('storyPage').classList.toggle('hidden',screen!=='workspace'||currentPage!=='story');$('flowPage').classList.toggle('hidden',screen!=='workspace'||currentPage!=='flow');if(screen==='manager')renderManager();else{if(currentPage==='review')renderReview();if(currentPage==='flow')renderFlow();if(currentPage==='story')renderStoryFrame();}}
function renderManager(){ $('subbar').classList.add('hidden');let p=currentParentItem();let filtered=projects.filter(x=>x.parent_id===currentParent&&x.module_type===currentModule);$('managerPage').innerHTML=`<div class="manager-layout"><aside class="manager-side"><div class="parent-create"><input id="newParentName" placeholder="新父项目名" onkeydown="if(event.key==='Enter')createParent()"><button class="btn" onclick="createParent()">新建父项目</button></div><div class="desc">先选择父级项目，再在右侧导入/打开子项目。</div>${parents.length?parents.map(item=>`<div class="parent-card ${item.parent_id===currentParent?'active':''}" onclick="currentParent='${esc(item.parent_id)}';render()"><b>${esc(item.parent_name||item.parent_id)}</b><div class="count-line">${parentCountLine(item)}</div><div class="count-line">${esc(item.updated_at||'')}</div><div class="actions"><button class="btn ghost small" onclick="event.stopPropagation();renameParent('${esc(item.parent_id)}')">改名</button><button class="btn danger small" onclick="event.stopPropagation();deleteParent('${esc(item.parent_id)}')">删除</button></div></div>`).join(''):'<div class="desc" style="margin-top:16px">暂无父项目，请新建或直接导入工程包。</div>'}</aside><section class="manager-main"><div class="manager-module-row"><button class="btn module-pill-btn ${currentModule==='material'?'':'ghost'}" onclick="currentModule='material';render()">美术</button><button class="btn module-pill-btn ${currentModule==='image'?'':'ghost'}" onclick="currentModule='image';render()">分镜</button><button class="btn module-pill-btn ${currentModule==='video'?'':'ghost'}" onclick="currentModule='video';render()">视频</button><button class="btn ghost manager-import" onclick="$('importFile').click()">导入工程包</button></div><div class="child-toolbar"><div><b>当前父级项目：</b>${p?esc(p.parent_name||p.parent_id):'未选择'}</div><span class="pill">${shortLabel(currentModule)} ${filtered.length}</span></div><div class="child-list">${p?filtered.map(item=>`<div class="child-card"><div class="child-main"><div class="child-title">${esc(item.child_name||item.project_name)}</div><div class="project-meta">${esc(item.module_label)} · ${statsText(item)}</div><div class="project-meta">导入：${esc(item.imported_at||'')} · 负责人：${esc(item.responsible_nickname||'-')} · 成本：$${Number(item.usage_usd||0).toFixed(6)}</div></div><div class="child-actions"><button class="btn small" onclick="openWorkspace('${item.workspace}')">打开审核</button><button class="btn ghost small" onclick="renameChild('${item.workspace}')">改名</button><button class="btn ghost small" onclick="location.href='/api/export_package?workspace=${encodeURIComponent(item.workspace)}'">导出审核包</button><button class="btn danger small" onclick="deleteWorkspace('${item.workspace}')">删除</button></div></div>`).join('')||'<div class="empty-center">当前模块下暂无工程包。导入任意模块工程包时，程序会自动按包内信息归位。</div>':'<div class="empty-center">请先选择或新建父项目。你也可以直接导入工程包，程序会自动创建对应父项目。</div>'}</div></section></div>`}
function statsText(p){let s=p.stats||{};if(p.module_type==='material')return `素材组/项：${s.unit_count||0} · 素材：${s.asset_count||0}`;return `SC：${s.scene_count||0} · 子栏：${s.shot_count||0} · 待选项：${s.candidate_count||0}`}
function emptyHtml(title){return `<div class="import-empty"><div class="import-card"><h1>${title}</h1><p class="desc">请从项目管理导入或打开工程。</p><button class="btn" onclick="showManager()">返回项目管理</button></div></div>`}
function renderReview(){syncNoteTextsBeforeRerender();if(!project){$('subbar').classList.add('hidden');$('reviewPage').innerHTML=emptyHtml(labelOf(currentModule)+'审核');return}if(moduleType==='material')renderMaterialReview();else renderSceneReview()}
function statusText(st){return {approved:'通过',rejected:'不通过',pending:'待审核'}[st||'pending']||'待审核'}function statusClass(st){return st==='approved'?'status-approved':st==='rejected'?'status-rejected':'status-pending'}function roundLabel(n){return {1:'一审',2:'二审',3:'三审'}[Number(n)||1]||'一审'}
function ensureReviewObj(rv){rv=rv||{};rv.rounds=rv.rounds||{};if(!rv.rounds['1'])rv.rounds['1']={label:'一审',status:'pending',locked:false,note:''};rv.current_round=Number(rv.current_round||1);rv.max_round_available=Number(rv.max_round_available||1);return rv}function ensureRoundObj(rv){rv=ensureReviewObj(rv);let k=String(rv.current_round||1);if(!rv.rounds[k])rv.rounds[k]={label:roundLabel(k),status:'pending',locked:false,note:''};return rv.rounds[k]}
function currentRound(rv){return ensureRoundObj(rv)}
function reviewControls(unitType,unitId,rv,noteId){rv=ensureReviewObj(rv);let r=currentRound(rv),opts='';for(let i=1;i<=Math.max(1,rv.max_round_available||1);i++)opts+=`<option value="${i}" ${i===rv.current_round?'selected':''}>${roundLabel(i)}</option>`;return `<div class="review-tools"><button class="btn small" onclick="submitReview('${unitType}','${unitId}','approved','${noteId}')">通过</button><button class="btn danger small" onclick="submitReview('${unitType}','${unitId}','rejected','${noteId}')">不通过</button><select id="round-${unitId.replace(/[^a-zA-Z0-9]/g,'_')}" onchange="changeRoundLocal('${unitType}','${unitId}',this.value)">${opts}</select><span class="status-label ${statusClass(r.status)}">${statusText(r.status)}</span>${r.locked?'<span class="locked-tag">已锁定</span>':''}</div>`}
function findAnyScene(id){let scs=[];if(moduleType==='material'){let wss=project.material_workspaces||{};['character','scene','object'].forEach(k=>{(wss[k]?.scenes||[]).forEach(s=>scs.push(s))})}scs=scs.concat(project.scenes||[]);return scs.find(s=>(s.scene_id||sceneCode(s))===id||sceneCode(s)===id)}
function findAnyShot(id){let scs=[];if(moduleType==='material'){let wss=project.material_workspaces||{};['character','scene','object'].forEach(k=>{(wss[k]?.scenes||[]).forEach(s=>scs.push(s))})}scs=scs.concat(project.scenes||[]);for(const s of scs){for(const sh of(s.shots||[])){if(sh.shot_id===id)return sh}}return null}
function editSceneTime(sceneId,which,part,val){let sc=findAnyScene(sceneId);if(!sc)return;let key=which==='start'?'time_start':'time_end';sc[key]=sc[key]||{h:0,m:0,s:0};sc[key][part]=Math.max(0,parseInt(val||0,10)||0);markReviewDirty()}
function editShotDuration(shotId,val){let sh=findAnyShot(shotId);if(!sh)return;sh.duration_seconds=Math.max(0,parseInt(val||0,10)||0);markReviewDirty()}
function markReviewDirty(){window.__reviewTimeDirty=true;let b=document.getElementById('reviewDirtyTip');if(b)b.style.display='inline'}
function changeRoundLocal(unitType,unitId,rn){if(moduleType==='material'&&(unitType==='scene'||unitType==='shot')){let catKey={'人物':'character','场景':'scene','道具':'object'};let key=catKey[materialCat]||'character';let scs=materialFlowScenes(key);if(unitType==='scene'){let sc=scs.find(s=>(s.scene_id||sceneCode(s))===unitId||sceneCode(s)===unitId);if(sc&&sc.review)sc.review.current_round=Number(rn)}else{scs.forEach(s=>(s.shots||[]).forEach(sh=>{if(sh.shot_id===unitId&&sh.review)sh.review.current_round=Number(rn)}))}renderReview();return}if(unitType==='material'){let rv=(project.material_review?.units||{})[unitId];if(rv)rv.current_round=Number(rn)}else{let sc=findScene(unitId);if(sc&&sc.review)sc.review.current_round=Number(rn)}renderReview()}
async function submitReview(unitType,unitId,status,noteId){let note=$(noteId)?.value||'';try{let r=await api('/api/review_submit',{method:'POST',body:{workspace,module_type:moduleType,unit_type:unitType,unit_id:unitId,status,note,reviewer:(window.deviceNick||'审核员')}});project=r.project;render();toast(status==='approved'?'已通过':'已标记不通过')}catch(e){toast(e.message)}}
function findScene(id){return (project.scenes||[]).find(s=>(s.scene_id||sceneCode(s))===id||sceneCode(s)===id)}function sceneNum(s){let n=Number(s.scene_number);if(n>0)return n;let m=String(s.scene_code||'').match(/(\d+)/);return m?Number(m[1]):999999}function sceneCode(s){let n=sceneNum(s);return n===999999?(s.scene_code||'SC-0000'):'SC-'+String(n).padStart(4,'0')}
function mediaPath(o){return o?.file_path||o?.path||o?.url||o?.src||o?.video_path||o?.image_path||o?.thumbnail_path||''}function fileUrl(p){p=String(p||'');if(!p)return '';if(/^data:|^https?:|^blob:/i.test(p))return p;p=p.replace(/^\/+/,'');if(p.startsWith('file/'))return '/'+p;return workspace?'/file/'+encodeURIComponent(workspace)+'/'+p:''}function isVideoPath(p){return /\.(mp4|mov|webm|mkv|avi)(\?|#|$)/i.test(String(p||''))}function isAudioPath(p){return /\.(mp3|wav|m4a|aac|ogg|flac)(\?|#|$)/i.test(String(p||''))}function mediaHtml(obj,kind='image',navShotId='',navIdx=0){let p=mediaPath(obj);let u=fileUrl(p);if(!p)return '<div class="empty">暂无文件</div>';if(kind==='audio'||isAudioPath(p))return `<div class="media-icon" ondblclick="openPreview('${u}','audio')">♫</div>`;if(kind==='video'||isVideoPath(p)){return `<video src="${u}" controls muted preload="metadata" onloadedmetadata="try{if(!this.dataset.seeked){this.currentTime=0.1;this.dataset.seeked=1}}catch(e){}" ondblclick="openPreview('${u}','video')"></video>`}return `<img src="${u}" ondblclick="${navShotId?("openCandidatePreview('"+jsstr(navShotId)+"',"+navIdx+",'"+u+"')"):("openPreview('"+u+"','image')")}">`}
function visibleCandidates(sh){return (sh.storyboard_candidates||[]).filter(c=>!c.review_deleted)}function moveCandidate(shotId,d){let arr=findShot(shotId)?.storyboard_candidates||[];let n=arr.filter(x=>!x.review_deleted).length;if(!n)return;candidateIndex[shotId]=((candidateIndex[shotId]||0)+d+n)%n;renderReview()}function findShot(id){for(const sc of project.scenes||[]){for(const sh of sc.shots||[]){if(sh.shot_id===id)return sh}}return null}
function renderSceneReview(){ $('subbar').classList.add('hidden');let scenes=[...(project.scenes||[])].sort((a,b)=>sceneNum(a)-sceneNum(b));$('reviewPage').innerHTML=scenes.map(sc=>{let sid=sc.scene_id||sceneCode(sc),rv=ensureReviewObj(sc.review);let noteId='note-'+sid;let shots=sc.shots||[];return `<section class="scene"><div class="scene-head"><div class="sc-code">${esc(sceneCode(sc))}</div><span class="pill">${shots.length} 个子栏</span>${sceneTimeEditor(sc)}${reviewControls('scene',sid,rv,noteId)}</div><div class="sc-content"><div class="strip">${shots.map(sh=>renderShotCard(sh)).join('')||'<div class="empty">暂无子栏</div>'}</div><div class="note-side">${noteSideHtml(sc,rv,'scene',sid,noteId)}</div></div></section>`}).join('')||'<div class="desc">工程中没有 SC。</div>'}
function noteTargetOf(unitType,unitId){return unitType==='shot'?findAnyShot(unitId):findAnyScene(unitId)}
function writeNoteText(unitType,unitId,roundNo,val){let t=noteTargetOf(unitType,unitId);if(!t)return;let rv=ensureReviewObj(t.review);let k=String(roundNo||rv.current_round||1);if(!rv.rounds[k])rv.rounds[k]={label:roundLabel(k),status:'pending',locked:false,note:''};rv.rounds[k].note=val;window.__reviewNoteDirty=true}
function setNoteText(unitType,unitId,val){let t=noteTargetOf(unitType,unitId);if(!t)return;let rv=ensureReviewObj(t.review);let r=ensureRoundObj(rv);r.note=val;window.__reviewNoteDirty=true}
function captureNoteTexts(){let map={};document.querySelectorAll('textarea.note-text[data-unit-type]').forEach(ta=>{map[ta.dataset.unitType+'|'+ta.dataset.unitId]={round:ta.dataset.round,val:ta.value}});return map}
function noteTextOf(texts,ut,uid){let e=texts&&texts[ut+'|'+uid];return (e&&typeof e==='object')?(e.val||''):(e||'')}
function applyNoteTexts(map){if(!map)return;Object.keys(map).forEach(k=>{let i=k.indexOf('|');let ut=k.slice(0,i),uid=k.slice(i+1);let e=map[k];let val=(e&&typeof e==='object')?e.val:e;let rnd=(e&&typeof e==='object')?e.round:null;writeNoteText(ut,uid,rnd,val)})}
function noteSideHtml(target,rv,unitType,unitId,noteId){let imgs=noteImagesOf(target,rv);let thumbs=imgs.map((im,i)=>`<div class="note-thumb"><img src="${esc(fileUrl(mediaPath(im)))}" onclick="openPreview('${jsstr(fileUrl(mediaPath(im)))}','image')"><button class="note-del" onclick="delNoteImage('${jsstr(unitType)}','${jsstr(unitId)}',${i})">×</button></div>`).join('');return `<div class="note-head">审核备注</div><div class="note-imgbox" ondragover="noteDragOver(event,this)" ondragleave="this.classList.remove('drag')" ondrop="dropOnNoteSide(event,'${jsstr(unitType)}','${jsstr(unitId)}',this)"><div class="note-imgs">${thumbs||'<span class="note-empty">拖入图片，或点＋备注图。支持本地图片、待选素材/分镜。</span>'}<button class="btn ghost small" onclick="addNoteImage('${jsstr(unitType)}','${jsstr(unitId)}')">＋备注图</button></div></div><textarea id="${noteId}" class="note-text" data-unit-type="${esc(unitType)}" data-unit-id="${esc(unitId)}" data-round="${(rv&&rv.current_round)||1}" placeholder="审核备注" oninput="writeNoteText('${jsstr(unitType)}','${jsstr(unitId)}',${(rv&&rv.current_round)||1},this.value)">${esc(currentRound(rv).note||'')}</textarea>`}
function noteDragOver(e,el){if(e.dataTransfer)e.dataTransfer.dropEffect='copy';e.preventDefault();el.classList.add('drag')}
async function dropOnNoteSide(e,unitType,unitId,el){e.preventDefault();el.classList.remove('drag');
  let texts=captureNoteTexts();
  // 本地文件
  if(e.dataTransfer.files&&e.dataTransfer.files.length){for(const f of Array.from(e.dataTransfer.files)){if(!/^image\//.test(f.type))continue;let rd=new FileReader();await new Promise(res=>{rd.onload=async()=>{try{let r=await api('/api/review_add_note_image',{method:'POST',body:{workspace,unit_type:unitType,unit_id:unitId,dataUrl:rd.result,round:curRoundOf(unitType,unitId),note_text:noteTextOf(texts,unitType,unitId)}});project=r.project}catch(err){toast('添加失败：'+err.message)}res()};rd.readAsDataURL(f)})}applyNoteTexts(texts);renderReview();toast('已添加备注图');return}
  // 待选素材/分镜：拖拽数据带 file url
  let url=e.dataTransfer.getData('note_image_url')||e.dataTransfer.getData('text/uri-list')||e.dataTransfer.getData('text/plain');
  if(url){try{let r=await api('/api/review_add_note_image_url',{method:'POST',body:{workspace,unit_type:unitType,unit_id:unitId,url:url,round:curRoundOf(unitType,unitId),note_text:noteTextOf(texts,unitType,unitId)}});project=r.project;applyNoteTexts(texts);renderReview();toast('已添加备注图')}catch(err){toast('添加失败：'+err.message)}}
}
function noteImagesOf(target,rv){let rn=String((rv&&rv.current_round)||1);let n=(target.review_notes_by_round||{})[rn]||{};return n.images||[]}
function noteImageStrip(target,rv,unitType,unitId){let imgs=noteImagesOf(target,rv);let thumbs=imgs.map((im,i)=>`<div class="note-thumb"><img src="${esc(fileUrl(mediaPath(im)))}" onclick="openPreview('${jsstr(fileUrl(mediaPath(im)))}','image')"><button class="note-del" onclick="delNoteImage('${jsstr(unitType)}','${jsstr(unitId)}',${i})">×</button></div>`).join('');return `<div class="note-imgs">${thumbs}<button class="btn ghost small" onclick="addNoteImage('${jsstr(unitType)}','${jsstr(unitId)}')">＋备注图</button></div>`}
async function addNoteImage(unitType,unitId){let inp=document.createElement('input');inp.type='file';inp.accept='image/*';inp.onchange=async()=>{let f=inp.files&&inp.files[0];if(!f)return;let texts=captureNoteTexts();let curText=noteTextOf(texts,unitType,unitId);let rd=new FileReader();rd.onload=async()=>{try{let r=await api('/api/review_add_note_image',{method:'POST',body:{workspace,unit_type:unitType,unit_id:unitId,dataUrl:rd.result,round:curRoundOf(unitType,unitId),note_text:curText}});project=r.project;applyNoteTexts(texts);renderReview();toast('已添加备注图')}catch(e){toast('添加失败：'+e.message)}};rd.readAsDataURL(f)};inp.click()}
async function delNoteImage(unitType,unitId,idx){let target=unitType==='shot'?findAnyShot(unitId):findAnyScene(unitId);if(!target)return;let texts=captureNoteTexts();let rn=String(curRoundOf(unitType,unitId));let slot=(target.review_notes_by_round||{})[rn];if(!slot||!slot.images)return;slot.images.splice(idx,1);try{await api('/api/review_save_note_images',{method:'POST',body:{workspace,unit_type:unitType,unit_id:unitId,round:rn,images:slot.images}});applyNoteTexts(texts);renderReview();toast('已删除')}catch(e){toast('删除失败：'+e.message)}}
function curRoundOf(unitType,unitId){let t=unitType==='shot'?findAnyShot(unitId):findAnyScene(unitId);let rv=ensureReviewObj(t&&t.review);return rv.current_round||1}
function sceneTimeEditor(sc){let sid=sc.scene_id||sceneCode(sc);let s=sc.time_start||{},e=sc.time_end||{};function inp(which,part,v){return `<input class="time-inp" type="number" min="0" value="${parseInt(v||0,10)||0}" onchange="editSceneTime('${jsstr(sid)}','${which}','${part}',this.value)">`}return `<span class="sc-time-edit">时间 ${inp('start','h',s.h)}时${inp('start','m',s.m)}分${inp('start','s',s.s)}秒 - ${inp('end','h',e.h)}时${inp('end','m',e.m)}分${inp('end','s',e.s)}秒</span>`}
function shotDurEditor(sh){return `<span class="sc-time-edit">时长 <input class="time-inp" type="number" min="0" value="${parseInt(sh.duration_seconds||sh.duration||0,10)||0}" onchange="editShotDuration('${jsstr(sh.shot_id)}',this.value)">秒</span>`}
function renderShotCard(sh){let cs=visibleCandidates(sh),idx=Math.min(candidateIndex[sh.shot_id]||0,Math.max(cs.length-1,0));candidateIndex[sh.shot_id]=idx;let c=cs[idx],kind=moduleType==='video'?'video':'image';let cu=c?fileUrl(mediaPath(c)):'';return `<div class="card"><div class="card-title">${esc(sh.shot_code||sh.shot_id||'子栏')} ${shotDurEditor(sh)}</div><div class="media" ${c?`draggable="true" ondragstart="noteCandDrag(event,'${jsstr(cu)}')"`:''}>${c?mediaHtml(c,kind,sh.shot_id,idx):'<div class="empty">暂无待选项</div>'}${c?`<button class="cand-del" title="删除该待选图" onclick="event.stopPropagation();delCandidate('${jsstr(sh.shot_id)}',${idx})">×</button>`:''}</div><div class="foot"><button class="arrow" onclick="moveCandidate('${sh.shot_id}',-1)">←</button><span>${cs.length?(idx+1)+' / '+cs.length:'0 / 0'}</span><button class="arrow" onclick="moveCandidate('${sh.shot_id}',1)">→</button></div></div>`}
function delCandidate(shotId,idx){let sh=(typeof findAnyShot==='function'?findAnyShot(shotId):null)||(typeof findShot==='function'?findShot(shotId):null);if(!sh)return;let cs=visibleCandidates(sh);let c=cs[idx];if(!c)return;if(!confirm('确定删除该待选图？删除后该图不再参与审核与制作。'))return;c.review_deleted=true;let left=visibleCandidates(sh);candidateIndex[shotId]=Math.min(idx,Math.max(left.length-1,0));renderReview();api('/api/review_delete_candidate',{method:'POST',body:{workspace,shot_id:shotId,index:idx}}).catch(e=>toast('删除未保存：'+e.message))}
function noteCandDrag(e,url){try{e.dataTransfer.setData('note_image_url',url);e.dataTransfer.setData('text/plain',url);e.dataTransfer.effectAllowed='copy'}catch(err){}}
function materialUnits(){return project?project.__material_units||[]:[]}function renderMaterialReview(){let tabs=['人物','场景','道具'].map(c=>`<button class="${materialCat===c?'active':''}" onclick="materialCat='${c}';renderReview()">${c}</button>`).join('');$('subbar').classList.remove('hidden');$('subbar').innerHTML=`<div class="cat-tabs">${tabs}</div><span class="pill">${labelOf(moduleType)} · ${materialCat}</span>`;
  let catKey={'人物':'character','场景':'scene','道具':'object'};let key=catKey[materialCat]||'character';
  let scenes=[...materialFlowScenes(key)].sort((a,b)=>sceneNum(a)-sceneNum(b));
  if(key==='object'){
    // 道具：单层子栏，平铺成卡片，每个子栏可审核
    let cards=[];scenes.forEach(sc=>(sc.shots||[]).forEach(sh=>{let sid=sh.shot_id,rv=ensureReviewObj(sh.review);let noteId='note-'+sid;cards.push(`<section class="scene"><div class="scene-head"><div class="sc-code">${esc(sh.shot_code||sh.shot_id||'道具')}</div>${reviewControls('shot',sid,rv,noteId)}</div><div class="sc-content"><div class="strip">${renderShotCard(sh)}</div><div class="note-side">${noteSideHtml(sh,rv,'shot',sid,noteId)}</div></div></section>`)}));
    $('reviewPage').innerHTML=cards.join('')||`<div class="desc">道具下暂无内容。</div>`;return;
  }
  $('reviewPage').innerHTML=scenes.map(sc=>{let sid=sc.scene_id||sceneCode(sc),rv=ensureReviewObj(sc.review);let noteId='note-'+sid;let shots=sc.shots||[];return `<section class="scene"><div class="scene-head"><div class="sc-code">${esc(sc.scene_code||sceneCode(sc))}</div><span class="pill">${shots.length} 个子栏</span>${sceneTimeEditor(sc)}${reviewControls('scene',sid,rv,noteId)}</div><div class="sc-content"><div class="strip">${shots.map(sh=>renderShotCard(sh)).join('')||'<div class="empty">暂无子栏</div>'}</div><div class="note-side">${noteSideHtml(sc,rv,'scene',sid,noteId)}</div></div></section>`}).join('')||`<div class="desc">${materialCat} 下暂无命名框。</div>`}
function renderMaterialUnit(u){let rv=ensureReviewObj((project.material_review?.units||{})[u.unit_id]);let noteId='note-'+u.unit_id.replace(/[^a-zA-Z0-9]/g,'_');return `<section class="unit"><div class="unit-head"><div class="unit-title">${esc(u.title)}</div><span class="pill">${u.unit_type==='group'?'素材组':'单个道具'} · ${u.assets.length} 个素材</span>${reviewControls('material',u.unit_id,rv,noteId)}</div><div class="strip">${u.assets.map(a=>`<div class="card"><div class="card-title">@${esc(a.name||'素材')}</div><div class="media">${mediaHtml(a,'image')}</div><div class="foot"><span class="meta">${esc(a.category||'素材')}</span></div></div>`).join('')||'<div class="empty">暂无素材</div>'}</div><div class="note-row"><textarea id="${noteId}" placeholder="审核备注">${esc(currentRound(rv).note||'')}</textarea></div></section>`}
function candDuration(c,sh){let vals=[c?.duration_seconds,c?.duration,c?.video_duration,c?.metadata?.duration,sh?.duration_seconds];for(let v of vals){let n=Number(v);if(isFinite(n)&&n>0)return n}return moduleType==='video'?5:2}
function buildStoryList(){storyList=[];storyIdx=0;storyTotalMs=0;storyCurrentMs=0;storyFrameOffsetMs=0;if(!project||moduleType==='material')return;let scenes=[...(project.scenes||[])].sort((a,b)=>sceneNum(a)-sceneNum(b));for(const sc of scenes){for(const sh of (sc.shots||[])){let cs=visibleCandidates(sh);if(!cs.length)continue;let idx=Math.min(Number(sh.review_selected_index||0),cs.length-1);let c=cs[idx];let p=mediaPath(c);if(!p)continue;let type=(moduleType==='video'||isVideoPath(p))?'video':'image';let dur= type==='image'?candDuration(c,sh):candDuration(c,{});storyList.push({type,src:fileUrl(p),duration:dur,info:`${sceneCode(sc)}/${sh.shot_code||''}`});}}let acc=0;for(const it of storyList){it.startMs=acc;it.endMs=acc+(Number(it.duration)||1)*1000;acc=it.endMs}storyTotalMs=acc}
function recomputeStoryTimeline(){let acc=0;for(const it of storyList){it.startMs=acc;it.endMs=acc+(Number(it.duration)||1)*1000;acc=it.endMs}storyTotalMs=acc}
function renderStoryFrame(){updateTop();if(moduleType==='material'){showPage('review');return}let item=storyList[storyIdx];if(!item){$('storyStage').innerHTML='<div class="story-black">暂无可播放内容</div>';updateStoryProgress();return}if(item.type==='video'){$('storyStage').innerHTML=`<video id="storyVideo" src="${item.src}" controls playsinline></video>`;let v=$('storyVideo');let offset=storyFrameOffsetMs||0;v.onloadedmetadata=()=>{if(isFinite(v.duration)&&v.duration>0){item.duration=v.duration;recomputeStoryTimeline();offset=Math.min(offset,Math.max(0,v.duration-.05));try{v.currentTime=offset}catch(e){}}updateStoryProgress();if(storyPlaying){v.play().catch(()=>{})}};v.onended=()=>{if(storyPlaying){storyIdx++;storyFrameOffsetMs=0;if(storyIdx>=storyList.length){storyIdx=Math.max(0,storyList.length-1);storyCurrentMs=storyTotalMs;pauseStory();updateStoryProgress();return}renderStoryFrame()}};v.ontimeupdate=updateStoryProgress;if(storyPlaying){v.play().catch(()=>{})}}else{$('storyStage').innerHTML=`<img src="${item.src}">`;storyFrameOffsetMs=Math.min(Math.max(storyFrameOffsetMs||0,0),(item.endMs||0)-(item.startMs||0));storyCurrentMs=(item.startMs||0)+storyFrameOffsetMs;updateStoryProgress()}}
function togglePlay(){storyPlaying?pauseStory():playStory()}function playStory(){if(!storyList.length)buildStoryList();if(!storyList.length)return;storyPlaying=true;$('playBtn').textContent='暂停';let item=storyList[storyIdx];if(item?.type==='video'){renderStoryFrame()}else{startImageFrame(storyFrameOffsetMs||0)}}function pauseStory(){storyPlaying=false;$('playBtn').textContent='播放';clearTimeout(storyTimer);let v=$('storyVideo');if(v)v.pause()}function restartStory(){pauseStory();storyIdx=0;storyFrameOffsetMs=0;renderStoryFrame()}function startImageFrame(offset=0){clearTimeout(storyTimer);let item=storyList[storyIdx];if(!item){pauseStory();return}if(item.type==='video'){storyFrameOffsetMs=offset;renderStoryFrame();return}storyFrameOffsetMs=offset;renderStoryFrame();let dur=(item.endMs||0)-(item.startMs||0);storyStartedAt=Date.now()-storyFrameOffsetMs;tickImageProgress();storyTimer=setTimeout(()=>{storyIdx++;storyFrameOffsetMs=0;if(storyIdx>=storyList.length){storyIdx=Math.max(0,storyList.length-1);pauseStory();storyCurrentMs=storyTotalMs;updateStoryProgress();return}if(storyPlaying)startImageFrame(0)},Math.max(0,dur-storyFrameOffsetMs))}function tickImageProgress(){if(!storyPlaying||moduleType!=='image')return;let item=storyList[storyIdx];storyCurrentMs=Math.min(item.endMs,(item.startMs||0)+(Date.now()-storyStartedAt));updateStoryProgress();requestAnimationFrame(tickImageProgress)}function updateStoryProgress(){let item=storyList[storyIdx];let cur=0;if(item){if(item.type==='video'){let v=$('storyVideo');cur=(item.startMs||0)+((v&&isFinite(v.currentTime))?v.currentTime*1000:(storyFrameOffsetMs||0))}else{cur=storyCurrentMs||((item.startMs||0)+(storyFrameOffsetMs||0))}}let pct=storyTotalMs?Math.min(100,Math.max(0,cur/storyTotalMs*100)):0;$('progressFill').style.width=pct+'%';$('playerInfo').textContent=`${fmt(cur)} / ${fmt(storyTotalMs)} · ${storyList.length?storyIdx+1:0} / ${storyList.length} ${item?.info||''}`}function fmt(ms){ms=Math.max(0,Math.floor((Number(ms)||0)/1000));let m=Math.floor(ms/60),s=ms%60;return String(m).padStart(2,'0')+':'+String(s).padStart(2,'0')}
function seekToMs(ms){if(!storyList.length)return;ms=Math.min(Math.max(0,ms),Math.max(storyTotalMs-1,0));let idx=storyList.findIndex(it=>ms>=it.startMs&&ms<it.endMs);if(idx<0)idx=storyList.length-1;let off=ms-(storyList[idx].startMs||0);let sameVideo=(idx===storyIdx)&&storyList[idx].type==='video'&&$('storyVideo');if(sameVideo){let v=$('storyVideo');storyFrameOffsetMs=off;try{v.currentTime=off/1000}catch(e){}updateStoryProgress();return}storyIdx=idx;storyFrameOffsetMs=off;if(storyList[idx].type==='video'){renderStoryFrame()}else{if(storyPlaying)startImageFrame(storyFrameOffsetMs);else renderStoryFrame()}}function seekStoryByEvent(e){let bar=$('storyProgress'),r=bar.getBoundingClientRect();let x=Math.min(Math.max(e.clientX-r.left,0),r.width);seekToMs(x/r.width*storyTotalMs)}function beginSeek(e){if(!storyList.length)buildStoryList();if(!storyList.length)return;e.preventDefault();seeking=true;wasPlayingBeforeSeek=storyPlaying;let v=$('storyVideo');if(v)v.pause();try{$('storyProgress').setPointerCapture(e.pointerId)}catch(_){}seekStoryByEvent(e);document.addEventListener('pointermove',dragSeek);document.addEventListener('pointerup',endSeek,{once:true})}function dragSeek(e){if(seeking){e.preventDefault();seekStoryByEvent(e)}}function endSeek(){if(!seeking)return;seeking=false;document.removeEventListener('pointermove',dragSeek);if(wasPlayingBeforeSeek){storyPlaying=true;$('playBtn').textContent='暂停';let item=storyList[storyIdx];if(item&&item.type==='video'){let v=$('storyVideo');if(v)v.play().catch(()=>{});else renderStoryFrame()}else{startImageFrame(storyFrameOffsetMs||0)}}}
function modelChips(obj){let keys=['image_model','video_model','model','model_id','gpt_model','openai_model','seedance_model','api_model','search_mode','image_ratio','video_ratio','image_resolution','video_resolution','duration','duration_seconds','resolution','quality','seed','audio','watermark'];let out=[];for(const [k,v] of Object.entries(obj||{})){if(keys.includes(k)||/model|模型|resolution|quality|ratio|duration|时长|search|seed|watermark|audio/i.test(k)){if(v!==undefined&&v!==null&&String(v)!=='')out.push(`<span class="model-chip">${esc(k)}: ${esc(v)}</span>`)}}return out.join('')||'<span class="desc">暂无模型字段</span>'}
function jsstr(s){return String(s??'').replace(/\\/g,'\\\\').replace(/'/g,"\\'").replace(/\r/g,'').replace(/\n/g,'\\n')}
function flowCategories(){return moduleType==='material'?['人物','场景','道具']:['人物','场景','道具','音频','视频']}
function flowTitle(){return moduleType==='video'?'生视频工作台':moduleType==='image'?'生分镜工作台':'美术素材工作台'}
function flowDesc(){if(moduleType==='video')return '每个 SC 下有多个子栏，每个子栏内部可切换标签页；右侧待选分镜空栏不属于标签页。审核端为只读镜像。';if(moduleType==='image')return '每个 SC 下有多个子栏，每个子栏内部可切换标签页；右侧待选分镜空栏不属于标签页。审核端为只读镜像。';return '按制作端主体工作台展示素材组与素材内容。审核端只读，保留分类和素材预览。'}
function assetById(id){return (project?.assets||[]).find(a=>a.asset_id===id||a.id===id)||null}
function flowAssetName(a){return String(a?.name||a?.title||a?.filename||'素材').replace(/^@+/,'')}
function flowMediaKind(obj){let p=mediaPath(obj);if(isAudioPath(p))return 'audio';if(isVideoPath(p))return 'video';return 'image'}
function flowMediaThumb(obj, extra=''){let p=mediaPath(obj),u=fileUrl(p),kind=flowMediaKind(obj);if(!p)return '<div class="media-icon">?</div>';if(kind==='audio')return `<div class="media-icon" ${extra} ondblclick="openPreview('${jsstr(u)}','audio')">♫</div>`;if(kind==='video'){let th=obj&&(obj.thumb_path||obj.thumbnail_path||obj.poster);let poster=th?fileUrl(th):'';return `<video src="${esc(u)}" ${poster?`poster="${esc(poster)}"`:''} preload="metadata" muted controls playsinline ${extra} onloadedmetadata="try{if(!this.dataset.seeked){this.currentTime=0.1;this.dataset.seeked=1}}catch(e){}" ondblclick="openPreview('${jsstr(u)}','video')"></video>`}return `<img src="${esc(u)}" ${extra} ondblclick="openPreview('${jsstr(u)}','image')">`}
function flowAssetCard(a){let u=fileUrl(mediaPath(a));return `<div class="asset-card" ondblclick="openPreview('${jsstr(u)}','${flowMediaKind(a)}')">${flowMediaThumb(a)}<div class="row"><div class="name" title="@${esc(flowAssetName(a))}">@${esc(flowAssetName(a))}</div></div><div class="meta" style="padding:0 6px 6px">${esc(a.category||'素材')}</div></div>`}
function flowRefChip(a){return `<div class="ref-chip">${flowMediaThumb(a)}<span title="@${esc(flowAssetName(a))}">@${esc(flowAssetName(a))}</span></div>`}
function flowGroupedCategory(cat){return cat==='人物'||cat==='场景'}
function flowGroupsFor(cat){let groups=(project?.asset_groups||[]).filter(g=>(g.category||'')===cat);let assets=(project?.assets||[]).filter(a=>!a.temporary&&(a.category||'')===cat);let used=new Set(groups.map(g=>g.group_id));let loose=assets.filter(a=>!a.group_id||!used.has(a.group_id));let arr=groups.map(g=>({...g,virtual:false}));if(loose.length)arr.push({group_id:'__loose_'+cat,category:cat,name:'未分组'+cat,virtual:true});return arr.sort((a,b)=>(a.sort_order||0)-(b.sort_order||0))}
function flowAssetsForGroup(cat,gid){return (project?.assets||[]).filter(a=>!a.temporary&&(a.category||'')===cat&&(gid&&gid.startsWith('__loose_')?(!a.group_id):a.group_id===gid))}
function flowSettingsValue(tab,keys,def='-'){let s=tab?.settings||{};for(const k of keys){let v=(s&&s[k]!==undefined?s[k]:tab?.[k]);if(v!==undefined&&v!==null&&String(v)!=='')return v}return def}
function flowActiveTab(shot){let tabs=shot.tabs||[];if(!tabs.length)return null;let id=flowTabState[shot.shot_id]||shot.active_tab_id||tabs[0].tab_id;return tabs.find(t=>t.tab_id===id)||tabs[0]}
function setFlowTab(shotId,tabId){flowTabState[shotId]=tabId;let m=document.querySelector('.workspace-main');let s=document.querySelector('.side-panel');let mt=m?m.scrollTop:0;let st=s?s.scrollTop:0;renderFlow();requestAnimationFrame(()=>{let m2=document.querySelector('.workspace-main');let s2=document.querySelector('.side-panel');if(m2)m2.scrollTop=mt;if(s2)s2.scrollTop=st})}
function flowTabLabel(t,i){return t?.tab_name||t?.title||t?.name||t?.label||('标签'+(i+1))}
function flowGenerated(tab){let arr=[];arr=arr.concat(tab?.generated_images||[]);arr=arr.concat(tab?.generated_videos||[]);arr=arr.concat(tab?.videos||[]);return arr}
function flowLastUserPrompt(tab){let msgs=tab?.messages||[];for(let i=msgs.length-1;i>=0;i--){if((msgs[i].role||'')==='user'&&msgs[i].content)return msgs[i].content}return tab?.draft_prompt||tab?.prompt||''}
function flowVideoInfo(tab){let last=tab?.last_video_status||{};if(last.state==='error')return `<div class="video-info error"><b>上次生成失败</b><br>${esc(last.message||'视频生成失败。')}</div>`;if(last.state==='success')return `<div class="video-info success"><b>生成完成</b><br>${esc(last.message||'视频生成完成。')}</div>`;return `<div class="video-info"><b>${moduleType==='video'?'视频':'图片'}生成信息</b><br>审核端只读展示制作端生成信息，不会调用模型或后端生成。</div>`}
function renderFlow(){if(!project){$('subbar').classList.add('hidden');$('flowPage').innerHTML=emptyHtml('制作流程');return}$('subbar').classList.add('hidden');if(moduleType==='material')renderMaterialFlow();else renderSceneFlow()}
function materialFlowScenes(key){let ws=(project&&project.material_workspaces&&project.material_workspaces[key])||null;let scs=ws&&Array.isArray(ws.scenes)?ws.scenes:[];if((!scs||!scs.length)&&key==='character'&&Array.isArray(project.scenes))scs=project.scenes;return scs}
function renderMaterialFlow(){
  let catKey={'人物':'character','场景':'scene','道具':'object'};
  let cat=materialCat||'人物'; let key=catKey[cat]||'character';
  let tabs=['人物','场景','道具'].map(c=>`<button class="tab-btn ${cat===c?'active':''}" onclick="materialCat='${c}';renderFlow()">${c}</button>`).join('');
  let scenes=[...materialFlowScenes(key)].sort((a,b)=>sceneNum(a)-sceneNum(b));
  let body;
  if(key==='object'){
    // 道具：单层子栏（每个 scene 的 shots 直接平铺，不显示父栏）
    let shots=[];scenes.forEach(sc=>(sc.shots||[]).forEach(sh=>shots.push([sc,sh])));
    body=`<section class="scene"><div class="scene-head"><span class="scene-title">道具</span><span class="pill">${shots.length} 个子栏</span><div class="scene-actions"><span class="status"><span class="dot green"></span>只读</span></div></div><div class="shot-list">${shots.map(([sc,sh])=>renderFlowShot(sc,sh)).join('')||'<div class="desc">暂无道具子栏</div>'}</div></section>`;
  } else {
    body=scenes.map(sc=>renderFlowScene(sc)).join('')||`<div class="desc">${cat} 下暂无命名框。</div>`;
  }
  $('flowPage').innerHTML=`<div class="flow-readonly"><main class="workspace-main"><div class="page-title">${flowTitle()}</div><div class="desc">${flowDesc()}</div><div class="asset-tabs" style="margin-bottom:12px">${tabs}</div><div id="flowSceneRoot">${body}</div></main><aside class="side-panel" id="flowSidePanel">${renderFlowSidePanel()}</aside></div>`;
}
function renderSceneFlow(){let scenes=[...(project.scenes||[])].sort((a,b)=>sceneNum(a)-sceneNum(b));$('flowPage').innerHTML=`<div class="flow-readonly"><main class="workspace-main"><div class="page-title">${flowTitle()}</div><div class="desc">${flowDesc()}</div><div id="flowSceneRoot">${scenes.map(sc=>renderFlowScene(sc)).join('')||'<div class="desc">暂无 SC。</div>'}</div></main><aside class="side-panel" id="flowSidePanel">${renderFlowSidePanel()}</aside></div>`}
function fmtSceneTime(sc){function pad(v){v=parseInt(v||0,10);return(isNaN(v)?0:v)}function one(t){t=t||{};return pad(t.h)+':'+String(pad(t.m)).padStart(2,'0')+':'+String(pad(t.s)).padStart(2,'0')}let s=sc.time_start||{},e=sc.time_end||{};let has=(s.h||s.m||s.s||e.h||e.m||e.s);return has?(one(s)+' - '+one(e)):''}
function renderFlowScene(sc){let shots=sc.shots||[];let mins=shots.reduce((s,sh)=>s+(Number(sh.duration_seconds||sh.duration||0)||0),0);let tr=fmtSceneTime(sc);return `<section class="scene"><div class="scene-head"><span class="scene-title">${esc(sceneCode(sc))}</span><span class="pill">${shots.length} 个子栏</span>${tr?`<span class="scene-time">时间 ${esc(tr)}</span>`:''}${mins?`<span class="scene-time">共 ${mins} 秒</span>`:''}<div class="scene-actions"><span class="status"><span class="dot green"></span>只读</span></div></div><div class="shot-list">${shots.map(sh=>renderFlowShot(sc,sh)).join('')||'<div class="desc">暂无子栏</div>'}</div></section>`}
function renderFlowShot(sc,sh){let tabs=sh.tabs||[];let tab=flowActiveTab(sh)||{};let idx=Math.max(0,tabs.indexOf(tab));let cands=visibleCandidates(sh);let status=cands.some(x=>x.review_status==='approved')?'<span class="dot green"></span>已确认':cands.length?'<span class="dot yellow"></span>待确认':'<span class="dot red"></span>未确认';return `<div class="shot"><div class="shot-head"><span class="handle">⋮⋮</span><span class="shot-title">${esc(sh.shot_code||sh.shot_id||'子栏')}</span>${sh.duration_seconds||sh.duration?`<span class="shot-duration">${esc(sh.duration_seconds||sh.duration)} 秒</span>`:''}<span class="status">${status}</span><div class="shot-actions"><span class="pill">只读</span></div></div><div class="shot-tabs">${tabs.map((t,i)=>`<button class="work-tab ${t===tab?'active':''}" onclick="setFlowTab('${jsstr(sh.shot_id)}','${jsstr(t.tab_id)}')"><span>${esc(flowTabLabel(t,i))}</span></button>`).join('')}${tabs.length?'<div class="tab-menu"><button class="btn ghost small disabled-action" title="审核端只读">＋</button><button class="btn ghost small disabled-action" title="审核端只读">复制</button></div>':''}</div><div class="shot-body"><div class="tabbed-work">${tab?renderFlowTabContent(sh,tab,idx):'<div class="desc">暂无标签页</div>'}</div>${renderFlowStoryPanel(sh)}</div></div>`}
function renderFlowTabContent(sh,tab,idx){return `<div class="tab-content"><div class="shot-left">${renderFlowSettings(tab,sh)}${renderFlowRefs(sh,tab)}</div><div class="shot-center">${flowVideoInfo(tab)}${renderFlowMessages(tab)}${renderFlowComposer(tab)}${renderFlowGenerated(tab)}</div></div>`}
function renderFlowSettings(tab,sh){if(moduleType==='material'){return `<div class="settings"><label>模型 <span class="ro-select">${esc(flowSettingsValue(tab,['image_model','model','model_id','gpt_model'],'-'))}</span></label><label>比例 <span class="ro-select">${esc(flowSettingsValue(tab,['image_ratio','ratio'],'-'))}</span></label><label>分辨率 <span class="ro-select">${esc(flowSettingsValue(tab,['image_resolution','resolution'],'-'))}</span></label><label>质量 <span class="ro-select">${esc(flowSettingsValue(tab,['quality','image_quality'],'-'))}</span></label><label>Seed <span class="ro-input">${esc(flowSettingsValue(tab,['seed'],'-'))}</span></label></div>`}return `<div class="settings"><label>模型 <span class="ro-select">${esc(flowSettingsValue(tab,['video_model','image_model','model','model_id'],'-'))}</span></label><label>比例 <span class="ro-select">${esc(flowSettingsValue(tab,['video_ratio','image_ratio','ratio'],'-'))}</span></label><label>分辨率 <span class="ro-select">${esc(flowSettingsValue(tab,['video_resolution','image_resolution','resolution'],'-'))}</span></label><label>时长 <span class="ro-input">${esc(flowSettingsValue(tab,['video_duration','duration_seconds','duration'],sh.duration_seconds||'-'))}</span></label><label>Seed <span class="ro-input">${esc(flowSettingsValue(tab,['seed'],'-'))}</span></label><label>同步音频 <span class="ro-select">${esc(flowSettingsValue(tab,['audio','sync_audio'],'-'))}</span></label><label>保留水印 <span class="ro-select">${esc(flowSettingsValue(tab,['watermark','keep_watermark'],'-'))}</span></label></div>`}
function renderFlowRefs(sh,tab){let refs=(tab.referenced_assets||[]).map(assetById).filter(Boolean);return `<div class="refs"><div class="refs-head"><span class="refs-title">当前标签页引用素材</span><button class="btn ghost small disabled-action" title="审核端只读">＋</button></div><div class="refs-list">${refs.map(flowRefChip).join('')||'<span class="desc">暂无引用素材。引用只属于当前标签页。</span>'}</div></div>`}
function renderFlowMessages(tab){let msgs=tab.messages||[];if(!msgs.length)return '';return msgs.slice(-20).map(m=>`<div class="msg ${esc(m.role||'message')}"><div class="msg-role">${(m.role||'')==='user'?'用户':'系统'} · ${esc(m.time||'')}</div><div class="bubble">${esc(m.content||'')}</div></div>`).join('')}
function renderFlowComposer(tab){let prompt=flowLastUserPrompt(tab);return `<div class="composer"><div class="prompt-area"><div class="prompt readonly">${esc(prompt||'')}</div>${prompt?'<div class="desc">提示词只读，可选中文本复制。</div>':''}</div><div class="composer-actions"><button class="btn small disabled-action" title="审核端不会调用模型">${moduleType==='video'?'生成视频':'生成图片'}</button></div></div>`}
function renderFlowGenerated(tab){let items=flowGenerated(tab);if(!items.length)return '';return `<div class="image-grid">${items.map(it=>`<div class="gen-img">${flowMediaThumb(it)}<div class="img-caption">${esc(it.name||it.title||it.image_id||it.video_id||'生成结果')}</div></div>`).join('')}</div>`}
function renderFlowStoryPanel(sh){let mode=sh.__flow_story_mode||'story';let active=mode==='notes';return `<div class="story"><div class="story-head"><div class="story-switch"><button class="tab-btn ${!active?'active':''}" onclick="let s=findShot('${jsstr(sh.shot_id)}');if(s){s.__flow_story_mode='story';renderFlow()}">待选分镜</button><button class="tab-btn ${active?'active':''}" onclick="let s=findShot('${jsstr(sh.shot_id)}');if(s){s.__flow_story_mode='notes';renderFlow()}">审核备注</button></div></div>${active?renderFlowReviewNote(sh):renderFlowStoryboardPanel(sh)}</div>`}
function renderFlowStoryboardPanel(sh){let cs=visibleCandidates(sh);let idx=flowCandidateIndex[sh.shot_id]||0;if(idx>=cs.length)idx=0;flowCandidateIndex[sh.shot_id]=idx;let c=cs[idx];let st=cs.some(x=>x.review_status==='approved')?'<span class="dot green"></span>已确认':cs.length?'<span class="dot yellow"></span>待确认':'';return `<div class="story-head">待选分镜素材 <span class="status">${st}</span><button class="btn ghost small disabled-action" title="审核端只读">＋素材</button></div><div class="story-preview">${c?flowMediaThumb(c):`<div class="story-empty">暂无待选素材<br>支持图片和视频<br>这是整个子栏共用区域</div>`}</div><div class="story-foot"><button class="arrow" onclick="moveFlowCandidate('${jsstr(sh.shot_id)}',-1)">←</button><span>${cs.length?(idx+1)+' / '+cs.length:'0 / 0'}</span><button class="arrow" onclick="moveFlowCandidate('${jsstr(sh.shot_id)}',1)">→</button></div>`}
function moveFlowCandidate(shotId,d){let sh=findShot(shotId);let n=visibleCandidates(sh||{}).length;if(!n)return;flowCandidateIndex[shotId]=((flowCandidateIndex[shotId]||0)+d+n)%n;renderFlow()}
function renderFlowReviewNote(sh){let rv=(findSceneByShotId(sh.shot_id)?.review)||{};let cr=currentRound(ensureReviewObj(rv));let imgs=[];let txt=cr.note||'';if(sh.review_notes_by_round){let rn=String(rv.current_round||1);let n=sh.review_notes_by_round[rn]||{};imgs=n.images||[];txt=n.text||txt}let idx=flowNoteIndex[sh.shot_id]||0;if(idx>=imgs.length)idx=0;let img=imgs[idx];return `<div class="story-head">审核备注</div><div class="note-preview">${img?flowMediaThumb(img):'<div class="story-empty">暂无审核备注素材</div>'}</div><div class="note-tools"><button class="arrow" onclick="moveFlowNote('${jsstr(sh.shot_id)}',-1)">←</button><span>${imgs.length?(idx+1)+' / '+imgs.length:'0 / 0'}</span><button class="arrow" onclick="moveFlowNote('${jsstr(sh.shot_id)}',1)">→</button></div><textarea class="note-box" readonly>${esc(txt||'')}</textarea>`}
function moveFlowNote(shotId,d){let sh=findShot(shotId);let imgs=[];if(sh?.review_notes_by_round){let scene=findSceneByShotId(shotId);let rn=String(scene?.review?.current_round||1);imgs=(sh.review_notes_by_round[rn]||{}).images||[]}let n=imgs.length;if(!n)return;flowNoteIndex[shotId]=((flowNoteIndex[shotId]||0)+d+n)%n;renderFlow()}
function findSceneByShotId(shotId){for(const sc of project?.scenes||[]){for(const sh of sc.shots||[]){if(sh.shot_id===shotId)return sc}}return null}
function renderFlowSidePanel(){return `<div class="side-tabs"><button id="flowSideAssets" class="${flowSideMode==='assets'?'active':''}" onclick="flowSideMode='assets';renderFlow()">参考素材</button><button id="flowSideBoard" class="${flowSideMode==='board'?'active':''}" onclick="flowSideMode='board';renderFlow()">${moduleType==='video'?'分镜视频':'分镜图片'}</button></div><div id="flowSideContent">${flowSideMode==='assets'?renderFlowSideAssets():renderFlowSideBoard()}</div>`}
function renderFlowSideAssets(){let cats=flowCategories();if(!cats.includes(flowSideCategory))flowSideCategory=cats[0];let tabs=`<div class="asset-tabs">${cats.map(c=>`<button class="tab-btn ${c===flowSideCategory?'active':''}" onclick="flowSideCategory='${jsstr(c)}';renderFlow()">${c}</button>`).join('')}</div>`;if(flowGroupedCategory(flowSideCategory)){let gs=flowGroupsFor(flowSideCategory);return tabs+`<div class="group-tools"><button class="btn small disabled-action" title="审核端只读">＋ 新建${flowSideCategory}</button><span class="desc mini-desc">组名只用于分类，@ 引用具体素材名</span></div><div class="asset-groups side-groups">${gs.map(g=>renderFlowAssetGroup(flowSideCategory,g,true)).join('')||`<div class="drop-hint">暂无${flowSideCategory}组</div>`}</div>`}let assets=(project?.assets||[]).filter(a=>!a.temporary&&(a.category||'')===flowSideCategory);return tabs+`<div class="asset-grid side-drop"><div class="drop-hint">当前分类：${esc(flowSideCategory)}（只读）</div>${assets.map(flowAssetCard).join('')||'<div class="drop-hint">暂无素材</div>'}</div>`}
function renderFlowAssetGroup(cat,g,compact=false){let items=flowAssetsForGroup(cat,g.group_id);let key=cat+'::'+g.group_id;let open=flowGroupOpen[key]!==false;return `<div class="asset-group ${open?'':'collapsed'}"><div class="asset-group-head"><div class="asset-group-top"><span class="handle group-handle">⋮⋮</span><b class="asset-group-title" title="${esc(g.name||'未命名组')}">${esc(g.name||'未命名组')}</b></div><div class="asset-group-controls"><button class="mini group-fold" onclick="toggleFlowGroup('${jsstr(key)}')">${open?'折叠':'展开'}</button><span class="meta asset-count">${items.length} 个素材</span></div></div>${open?`<div class="group-grid asset-grid">${items.map(flowAssetCard).join('')||'<div class="desc">暂无素材。</div>'}</div>`:''}</div>`}
function toggleFlowGroup(key){flowGroupOpen[key]=flowGroupOpen[key]===false?true:false;renderFlow()}
function renderFlowSideBoard(){let cards=[];(project?.scenes||[]).forEach(sc=>(sc.shots||[]).forEach(sh=>{let cs=visibleCandidates(sh);let idx=flowCandidateIndex[sh.shot_id]||0;if(idx>=cs.length)idx=0;let c=cs[idx];cards.push(`<div class="board-card"><h4>${esc(sceneCode(sc))} / ${esc(sh.shot_code||sh.shot_id||'子栏')}</h4><div class="board-img">${c?flowMediaThumb(c):'<span class="desc">暂无待选素材</span>'}</div><div class="board-foot"><button class="arrow" onclick="moveFlowCandidate('${jsstr(sh.shot_id)}',-1)">←</button><span>${cs.length?(idx+1)+' / '+cs.length:'0 / 0'}</span><button class="arrow" onclick="moveFlowCandidate('${jsstr(sh.shot_id)}',1)">→</button></div></div>`)}));return `<div class="board-grid">${cards.join('')||'<div class="desc">暂无分镜素材</div>'}</div>`}
function openPreview(src,type,list,index){if(!src)return;let img=$('previewImg'),v=$('previewVideo'),a=$('previewAudio');img.classList.add('hidden');v.classList.add('hidden');if(a)a.classList.add('hidden');let arr;if(Array.isArray(list))arr=list;else if(typeof list==='string'&&__pvGroups[list])arr=__pvGroups[list];else arr=[src];__pvList=arr.filter(Boolean);let _i=(typeof index==='number')?index:__pvList.findIndex(it=>__pvSrc(it)===src);__pvIndex=(_i>=0?_i:0);if(type==='audio'||isAudioPath(src)){if(a){a.classList.remove('hidden');a.src=src;a.play().catch(()=>{})}}else if(type==='video'||isVideoPath(src)){v.classList.remove('hidden');v.src=src;v.play().catch(()=>{})}else{img.classList.remove('hidden');img.src=src;setupPreviewZoom(img)}$('preview').classList.add('show');__pvSyncDelBtn()}
let __pvList=[],__pvIndex=0,__pvGroups={};
function pvGroup(k,a){__pvGroups[k]=(a||[]).filter(Boolean);return k}
function __pvSrc(it){return (it&&typeof it==='object')?it.src:it}
function reviewScenesInView(){if(typeof moduleType!=='undefined'&&moduleType==='material'){const ck={'人物':'character','场景':'scene','道具':'object'};try{return materialFlowScenes(ck[materialCat]||'character')||[]}catch(e){return []}}return [...(project&&project.scenes||[])].sort((a,b)=>sceneNum(a)-sceneNum(b))}
function reviewImageItems(){let items=[];reviewScenesInView().forEach(sc=>{(sc.shots||[]).forEach(sh=>{visibleCandidates(sh).forEach((c,i)=>{items.push({src:fileUrl(mediaPath(c)),shotId:sh.shot_id,idx:i})})})});return items}
function openCandidatePreview(shotId,idx,src){let items=reviewImageItems();let pos=items.findIndex(it=>it.shotId===shotId&&it.idx===idx);if(pos<0)pos=items.findIndex(it=>it.src===src);if(pos<0){items=[{src:src,shotId:shotId,idx:idx}];pos=0}openPreview(src,'image',items,pos)}
function __pvSyncDelBtn(){const md=$('preview');const btn=$('pvDel');if(!md||!btn)return;const it=__pvList[__pvIndex];const can=!!(it&&typeof it==='object'&&it.shotId!==undefined&&it.shotId!==null);md.classList.toggle('has-del',can)}
function __pvNav(d){if(!__pvList||__pvList.length<2)return;__pvIndex=(__pvIndex+d+__pvList.length)%__pvList.length;let im=$('previewImg');if(im){im.src=__pvSrc(__pvList[__pvIndex]);__pvZoom.scale=1;__pvZoom.tx=0;__pvZoom.ty=0;__pvApply(im)}__pvSyncDelBtn()}
function delPreviewCurrent(){let it=__pvList[__pvIndex];if(!it||typeof it!=='object'||it.shotId===undefined)return;if(!confirm('确定删除该待选图？删除后不再参与审核与制作。'))return;let sh=(typeof findAnyShot==='function'?findAnyShot(it.shotId):null)||(typeof findShot==='function'?findShot(it.shotId):null);if(sh){let cs=visibleCandidates(sh);let c=cs[it.idx];if(c)c.review_deleted=true}api('/api/review_delete_candidate',{method:'POST',body:{workspace:workspace,shot_id:it.shotId,index:it.idx}}).catch(e=>toast('删除未保存：'+e.message));renderReview();let items=reviewImageItems();if(!items.length){$('preview').classList.remove('show');return}__pvList=items;__pvIndex=Math.min(__pvIndex,items.length-1);let im=$('previewImg');if(im){im.src=__pvSrc(__pvList[__pvIndex]);__pvZoom.scale=1;__pvZoom.tx=0;__pvZoom.ty=0;__pvApply(im)}__pvSyncDelBtn()}
let __pvZoom={scale:1,tx:0,ty:0,dragging:false,sx:0,sy:0,bound:false};
function __pvApply(img){img.style.transform='translate('+__pvZoom.tx+'px,'+__pvZoom.ty+'px) scale('+__pvZoom.scale+')';img.style.cursor=__pvZoom.scale>1?'grab':'default'}
function setupPreviewZoom(img){
  __pvZoom.scale=1;__pvZoom.tx=0;__pvZoom.ty=0;__pvZoom.dragging=false;__pvApply(img);
  img.style.transformOrigin='0 0';img.style.transition='transform .05s linear';
  if(__pvZoom.bound)return;__pvZoom.bound=true;
  // 阻止点图片时关闭弹窗
  img.addEventListener('click',e=>e.stopPropagation());
  img.addEventListener('dblclick',e=>{e.stopPropagation();__pvZoom.scale=1;__pvZoom.tx=0;__pvZoom.ty=0;__pvApply(img)});
  img.addEventListener('wheel',e=>{e.preventDefault();e.stopPropagation();const r=img.getBoundingClientRect();const s=__pvZoom.scale;let ns=Math.min(8,Math.max(1,s*(e.deltaY<0?1.15:1/1.15)));if(ns===s)return;__pvZoom.tx=__pvZoom.tx+(e.clientX-r.left)*(1-ns/s);__pvZoom.ty=__pvZoom.ty+(e.clientY-r.top)*(1-ns/s);__pvZoom.scale=ns;if(ns<=1){__pvZoom.tx=0;__pvZoom.ty=0}__pvApply(img)},{passive:false});
  window.addEventListener('keydown',e=>{const md=$('preview');if(!md||!md.classList.contains('show'))return;let im=$('previewImg');if(!im||im.classList.contains('hidden'))return;if(e.key==='ArrowLeft'){__pvNav(-1);e.preventDefault()}else if(e.key==='ArrowRight'){__pvNav(1);e.preventDefault()}});
  img.addEventListener('mousedown',e=>{if(__pvZoom.scale<=1)return;e.preventDefault();e.stopPropagation();__pvZoom.dragging=true;__pvZoom.sx=e.clientX-__pvZoom.tx;__pvZoom.sy=e.clientY-__pvZoom.ty;img.style.cursor='grabbing'});
  window.addEventListener('mousemove',e=>{if(!__pvZoom.dragging)return;__pvZoom.tx=e.clientX-__pvZoom.sx;__pvZoom.ty=e.clientY-__pvZoom.sy;let im=$('previewImg');if(im)__pvApply(im)});
  window.addEventListener('mouseup',()=>{if(__pvZoom.dragging){__pvZoom.dragging=false;let im=$('previewImg');if(im)im.style.cursor=__pvZoom.scale>1?'grab':'default'}});
}function stopPreview(){let v=$('previewVideo');v.pause();v.removeAttribute('src');let a=$('previewAudio');if(a){a.pause();a.removeAttribute('src')}__pvZoom.scale=1;__pvZoom.tx=0;__pvZoom.ty=0;let im=$('previewImg');if(im){im.style.transform='';im.style.cursor='default'}}
function exportExcel(){if(!workspace){alert('请先打开一个审核工程');return}location.href='/api/export_excel?workspace='+encodeURIComponent(workspace)}function exportPackage(){if(!workspace){alert('请先打开一个审核工程');return}location.href='/api/export_package?workspace='+encodeURIComponent(workspace)}

/* ===== 数据中心接入（V7） ===== */
window.deviceNick = localStorage.getItem('dilanDeviceNickname') || '';
let dcDownloadData=null, dcConflictToken=null;
function dcInjectDom(){
  if(document.getElementById('dcDownloadModal')) return;
  const wrap=document.createElement('div');
  wrap.innerHTML =
    '<div class="modal" id="nickModal"><div style="width:min(420px,92vw);background:#0f141a;border:1px solid rgba(255,255,255,.14);border-radius:18px;padding:22px">'+
    '<h2 style="margin:0 0 8px">第一次使用，请填写昵称</h2>'+
    '<p style="color:#8a9098;font-size:13px;line-height:1.7">这个昵称用于标识这台设备的使用者，会作为上传到数据中心项目的“生成者”，并写入审核表单的生成者列。</p>'+
    '<input id="nickInput" placeholder="例如：张三 / 审核01" style="width:100%;margin:10px 0 14px;padding:9px">'+
    '<div style="display:flex;justify-content:flex-end"><button class="btn" onclick="dcSaveNick()">确认进入</button></div></div></div>'+
    '<div class="modal" id="dcDownloadModal"><div class="dc-card">'+
    '<div class="dc-card-h"><h2 style="margin:0">数据中心 · 下载项目</h2><div>'+
    '<button class="btn ghost small" onclick="dcOpenDownload()">刷新</button> '+
    '<button class="btn ghost small" onclick="dcCloseDownload()">关闭</button></div></div>'+
    '<div id="dcDownloadBody" class="dc-body"></div></div></div>'+
    '<div class="modal" id="dcConflictModal"><div style="width:min(460px,92vw);background:#0f141a;border:1px solid rgba(255,255,255,.14);border-radius:18px;padding:22px">'+
    '<h2 style="margin:0 0 8px">数据中心已存在同名项目</h2><p id="dcConflictDesc" style="color:#8a9098;font-size:13px;line-height:1.7"></p>'+
    '<div style="display:flex;justify-content:flex-end;gap:8px;margin-top:14px">'+
    '<button class="btn ghost" onclick="dcConflictCancel()">取消</button>'+
    '<button class="btn ghost" onclick="dcResolve(\'new_version\')">生成新版本</button>'+
    '<button class="btn danger" onclick="dcResolve(\'overwrite\')">覆盖已有</button></div></div></div>';
  document.body.appendChild(wrap);
  const st=document.createElement('style');
  st.textContent='.dc-card{width:min(760px,94vw);max-height:86vh;display:flex;flex-direction:column;background:#0f141a;border:1px solid rgba(255,255,255,.12);border-radius:18px;padding:18px}'+
    '.dc-card-h{display:flex;align-items:center;justify-content:space-between;margin-bottom:10px}'+
    '.dc-body{overflow:auto;flex:1;min-height:120px}'+
    '.dc-parent{border:1px solid rgba(255,255,255,.1);background:#0b0f14;border-radius:14px;padding:12px;margin-bottom:12px}'+
    '.dc-parent-h{display:flex;align-items:center;gap:10px;flex-wrap:wrap;margin-bottom:8px}'+
    '.dc-count{font-size:12px;color:#8a9098}.dc-parent-h .btn{margin-left:auto}'+
    '.dc-mod{margin:8px 0}.dc-mod-h{font-weight:900;color:#00e5a0;font-size:13px;margin-bottom:6px}'+
    '.dc-child{display:flex;align-items:center;gap:10px;background:#111720;border:1px solid rgba(255,255,255,.08);border-radius:10px;padding:8px 10px;margin-bottom:6px}'+
    '.dc-child>div:first-child{flex:1;min-width:0}.dc-child b{font-size:13px}'+
    '.dc-meta{font-size:11px;color:#8a9098;margin-top:2px}'+
    '.dc-empty{color:#8a9098;text-align:center;padding:36px 10px;font-size:13px;line-height:1.7}';
  document.head.appendChild(st);
  if(!window.deviceNick){ document.getElementById('nickModal').classList.add('show'); }
}
function dcSaveNick(){const v=(document.getElementById('nickInput').value||'').trim(); if(!v)return toast('请填写昵称'); localStorage.setItem('dilanDeviceNickname',v); window.deviceNick=v; document.getElementById('nickModal').classList.remove('show'); toast('已记录昵称：'+v)}
async function dcConfig(){let c={url:''};try{c=await api('/api/datacenter/config')}catch(e){} const url=prompt('数据中心地址（局域网，例如 http://192.168.1.99:8777）：',c.url||''); if(url===null)return; try{await api('/api/datacenter/config/save',{method:'POST',body:{url:url.trim()}});toast('数据中心地址已保存')}catch(e){toast('保存失败：'+e.message)}}
async function saveReviewTimes(){
  if(!workspace) return;
  // 收集所有 scene 时间 + shot 时长（含美术三套 workspace）
  let scenes=[];
  if(moduleType==='material'){let wss=project.material_workspaces||{};['character','scene','object'].forEach(k=>{(wss[k]?.scenes||[]).forEach(s=>scenes.push(s))})}
  scenes=scenes.concat(project.scenes||[]);
  let times=scenes.map(sc=>({scene_id:sc.scene_id||sceneCode(sc),time_start:sc.time_start||null,time_end:sc.time_end||null,
    shots:(sc.shots||[]).map(sh=>({shot_id:sh.shot_id,duration_seconds:sh.duration_seconds}))}));
  await api('/api/review_save_times',{method:'POST',body:{workspace,times}});
  window.__reviewTimeDirty=false;let b=document.getElementById('reviewDirtyTip');if(b)b.style.display='none';
}
async function dcUploadProject(){
  if(!workspace) return toast('请先打开一个审核工程');
  if(!window.deviceNick){ document.getElementById('nickModal').classList.add('show'); return toast('请先填写昵称'); }
  if(window.__uploading){ return toast('正在上传中，请稍候…'); }
  window.__uploading=true; showUploadStatus('正在上传到数据中心…');
  try{ if(window.__reviewTimeDirty){await saveReviewTimes()}
    const r=await api('/api/datacenter/upload',{method:'POST',body:{workspace,user_name:window.deviceNick,on_conflict:'ask'}});
    if(r.conflict){ window.__uploading=false; hideUploadStatus(); dcConflictToken=r.token; document.getElementById('dcConflictDesc').textContent='数据中心已存在「'+r.incoming.module_label+' / '+r.incoming.child_name+'」。覆盖会替换原工程包（底层保留备份）；生成新版本会保留两者。'; document.getElementById('dcConflictModal').classList.add('show'); return; }
    hideUploadStatus(); toast('上传完成：'+((r.child&&r.child.child_name)||''));
  }catch(e){ hideUploadStatus(); alert('上传失败：'+e.message); toast('上传失败：'+e.message); }
  finally{ window.__uploading=false; }
}
function showUploadStatus(msg){let b=document.getElementById('uploadStatusBar');if(!b){b=document.createElement('div');b.id='uploadStatusBar';b.style.cssText='position:fixed;left:50%;top:14px;transform:translateX(-50%);z-index:6000;background:#11321f;border:1px solid #00e5a0;color:#aef5d6;padding:10px 18px;border-radius:10px;font-weight:700;box-shadow:0 6px 24px rgba(0,0,0,.4)';document.body.appendChild(b)}b.textContent=msg;b.style.display='block'}
function hideUploadStatus(){let b=document.getElementById('uploadStatusBar');if(b)b.style.display='none'}
function dcConflictCancel(){dcConflictToken=null;document.getElementById('dcConflictModal').classList.remove('show')}
async function dcResolve(action){ if(!dcConflictToken){dcConflictCancel();return} try{await api('/api/datacenter/resolve-conflict',{method:'POST',body:{token:dcConflictToken,action}});toast(action==='overwrite'?'已覆盖':'已生成新版本')}catch(e){toast('处理失败：'+e.message)} dcConflictToken=null;document.getElementById('dcConflictModal').classList.remove('show')}
async function dcOpenDownload(){
  document.getElementById('dcDownloadModal').classList.add('show');
  document.getElementById('dcDownloadBody').innerHTML='<div class="dc-empty">正在读取数据中心…</div>';
  try{ dcDownloadData=await api('/api/datacenter/projects'); dcRenderDownload(); }
  catch(e){ document.getElementById('dcDownloadBody').innerHTML='<div class="dc-empty">读取数据中心失败：'+esc(e.message)+'<br>请检查“数据中心地址”是否正确、服务端是否已启动、是否在同一局域网。</div>'; }
}
function dcCloseDownload(){document.getElementById('dcDownloadModal').classList.remove('show')}
function dcRenderDownload(){
  const r=dcDownloadData; if(!r)return;
  const labels=r.module_labels||{material:'美术',image:'分镜',video:'视频'};
  const order=r.module_order||['material','image','video'];
  const body=document.getElementById('dcDownloadBody');
  if(!r.parents||!r.parents.length){body.innerHTML='<div class="dc-empty">数据中心暂无项目。</div>';return}
  body.innerHTML=r.parents.map(p=>{
    const c=p.counts||{};
    const groups=order.map(mt=>{
      const kids=(p.children||[]).filter(x=>x.module_type===mt);
      if(!kids.length)return '';
      return '<div class="dc-mod"><div class="dc-mod-h">'+esc(labels[mt]||mt)+'（'+kids.length+'）</div>'+
        kids.map(k=>{const gens=(k.generator_names&&k.generator_names.length)?k.generator_names.join('、'):(k.generator_name||'-');
          return '<div class="dc-child"><div><b>'+esc(k.child_name)+'</b><div class="dc-meta">生成者：'+esc(gens)+' · 来源：'+esc(k.source_label||'-')+' · '+esc(k.uploaded_at||'')+'</div></div>'+
            '<button class="btn small" onclick="dcDownloadChild(\''+esc(k.child_id)+'\',\''+esc(k.child_name)+'\')">下载</button></div>';
        }).join('')+'</div>';
    }).join('');
    return '<div class="dc-parent"><div class="dc-parent-h"><b>'+esc(p.parent_name)+'</b>'+
      '<span class="dc-count">美术 '+(c.material||0)+' · 分镜 '+(c.image||0)+' · 视频 '+(c.video||0)+'</span>'+
      '<button class="btn ghost small" onclick="dcDownloadParent(\''+esc(p.parent_id)+'\',\''+esc(p.parent_name)+'\')">下载整个父项目</button></div>'+
      (groups||'<div class="dc-meta">该父项目暂无子项目</div>')+'</div>';
  }).join('');
}
async function dcDownloadChild(cid,name){ try{ toast('正在下载：'+name);
  const r=await api('/api/datacenter/download-child',{method:'POST',body:{child_id:cid,name:name}});
  toast('已下载并归位：'+(r.child_name||name)); await loadProjects();
}catch(e){ toast('下载失败：'+e.message); } }
async function dcDownloadParent(pid,pname){ try{ toast('正在下载父项目：'+pname);
  const r=await api('/api/datacenter/download-parent',{method:'POST',body:{parent_id:pid}});
  toast('父项目【'+pname+'】下载完成，归位 '+(r.imported||0)+'/'+(r.total||0)+' 个子项目'); await loadProjects();
}catch(e){ toast('下载父项目失败：'+e.message); } }
dcInjectDom();

/* 数据中心页面复刻接入（审核端） */
/* ===== 数据中心页面复刻（嵌入制作端/审核端，数据走本端 /api/datacenter/* 代理） ===== */
(function(){
  if(window.__dcReplicaInstalled) return; window.__dcReplicaInstalled=true;
  var DATA={parents:[],module_labels:{material:'美术',image:'分镜',video:'视频'},module_order:['material','image','video']};
  var curParent=null, curModule='material', curSource='creator', mMode=false, mSel=[];
  function E(s){return String(s==null?'':s).replace(/[&<>"]/g,function(c){return({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c])})}
  function fmtSize(n){n=Number(n||0);if(n<1024)return n+' B';if(n<1048576)return(n/1024).toFixed(1)+' KB';if(n<1073741824)return(n/1048576).toFixed(1)+' MB';return(n/1073741824).toFixed(2)+' GB'}
  function dcToast(m){ if(typeof toast==='function') toast(m); else alert(m); }
  async function dcApi(path,opts){
    var r=await fetch(path,opts); var t=await r.text(); var d; try{d=t?JSON.parse(t):{}}catch(e){d={error:t}}
    if(!r.ok && !d.conflict) throw new Error(d.error||('HTTP '+r.status)); return d;
  }
  function srcMatch(c){var ct=c.client_type||'';return curSource==='reviewer'?ct==='reviewer':ct!=='reviewer'}
  function pCounts(p){var o={material:0,image:0,video:0};(p.children||[]).forEach(function(c){if(srcMatch(c)&&o[c.module_type]!=null)o[c.module_type]++});return o}
  function curParentItem(){return DATA.parents.find(function(p){return p.parent_id===curParent})||null}

  function ensureDom(){
    if(document.getElementById('dcReplicaOverlay')) return;
    var ov=document.createElement('div'); ov.id='dcReplicaOverlay'; ov.className='dcr-overlay';
    ov.innerHTML=
      '<div class="dcr-app">'+
      '  <header class="dcr-top">'+
      '    <div class="dcr-brand"><i>蓝</i><span>数据中心 · 下载项目</span></div>'+
      '    <span class="dcr-pill" id="dcrHostTag"></span>'+
      '    <div style="flex:1"></div>'+
      '    <button class="btn ghost" onclick="window.__dcReplica.refresh()">刷新</button>'+
      '    <button class="btn ghost" onclick="window.__dcReplica.close()">关闭</button>'+
      '  </header>'+
      '  <div class="dcr-body">'+
      '    <aside class="dcr-parent-pane"><div class="dcr-pane-title">父项目</div>'+
      '      <div class="dcr-desc">点击父项目查看其下子项目；可整包下载整个父项目。</div>'+
      '      <div id="dcrParentList" class="dcr-parent-list"></div></aside>'+
      '    <section class="dcr-child-pane">'+
      '      <div class="dcr-module-bar"><div class="dcr-module-list" id="dcrModuleList"></div>'+
      '        <div class="dcr-module-actions">'+
      '          <button class="btn ghost" id="dcrMergeToggle" onclick="window.__dcReplica.toggleMerge()">合并子项目</button>'+
      '          <button class="btn" id="dcrMergeRun" style="display:none" onclick="window.__dcReplica.runMerge()">执行合并(0)</button>'+
      '        </div></div>'+
      '      <div class="dcr-child-toolbar"><div id="dcrParentTitle"><b>当前父项目：</b>未选择</div>'+
      '        <span class="dcr-pill" id="dcrModuleCount">—</span>'+
      '        <div class="dcr-source-switch" id="dcrSourceSwitch">'+
      '          <button class="dcr-src-btn" data-src="creator" onclick="window.__dcReplica.source(\'creator\')">制作端</button>'+
      '          <button class="dcr-src-btn" data-src="reviewer" onclick="window.__dcReplica.source(\'reviewer\')">审核端</button>'+
      '        </div></div>'+
      '      <div id="dcrChildList" class="dcr-child-list"></div>'+
      '    </section>'+
      '  </div>'+
      '</div>'+
      '<div class="dcr-preview" id="dcrPreview" onclick="this.classList.remove(\'show\')"><img id="dcrPreviewImg" alt="预览"></div>';
    document.body.appendChild(ov);
    var st=document.createElement('style');
    st.textContent=
      '.dcr-overlay{position:fixed;inset:0;z-index:3000;background:#0a0c0f;display:none}'+
      '.dcr-overlay.show{display:block}'+
      '.dcr-app{height:100%;display:flex;flex-direction:column;color:#e8eaed;font-family:"Microsoft YaHei",Arial,sans-serif}'+
      '.dcr-top{height:52px;display:flex;align-items:center;gap:10px;padding:0 14px;border-bottom:1px solid rgba(255,255,255,.08);background:#090b0e;flex-shrink:0}'+
      '.dcr-brand{font-weight:900;display:flex;align-items:center;gap:8px;font-size:16px}'+
      '.dcr-brand i{width:26px;height:26px;border-radius:8px;background:#00e5a0;color:#00120c;display:flex;align-items:center;justify-content:center;font-style:normal;font-size:15px}'+
      '.dcr-pill{background:#14191f;border:1px solid rgba(255,255,255,.1);border-radius:9px;padding:5px 9px;font-size:12px;color:#8a9098}'+
      '.dcr-body{flex:1;display:flex;min-height:0}'+
      '.dcr-parent-pane{width:330px;flex-shrink:0;background:#0d1116;border-right:1px solid rgba(255,255,255,.08);padding:18px;overflow:auto}'+
      '.dcr-child-pane{flex:1;min-width:0;background:#0b1015;overflow:auto;display:flex;flex-direction:column}'+
      '.dcr-pane-title{font-size:18px;font-weight:900;margin:0 0 4px}'+
      '.dcr-desc{font-size:12px;color:#8a9098;margin-bottom:14px;line-height:1.6}'+
      '.dcr-parent-list{display:flex;flex-direction:column;gap:10px}'+
      '.dcr-parent-card{border:1px solid rgba(255,255,255,.09);background:#111820;border-radius:14px;padding:12px;cursor:pointer}'+
      '.dcr-parent-card:hover{border-color:rgba(0,229,160,.35)}'+
      '.dcr-parent-card.active{border-color:rgba(0,229,160,.55);background:rgba(0,229,160,.07)}'+
      '.dcr-parent-card .pn{font-size:15px;font-weight:900;margin-bottom:6px}'+
      '.dcr-parent-card .pc{font-size:12px;color:#8a9098;margin-bottom:3px}'+
      '.dcr-parent-card .pt{font-size:11px;color:#596271;margin-bottom:9px}'+
      '.dcr-module-bar{display:flex;align-items:center;min-height:70px;padding:14px 22px;border-bottom:1px solid rgba(255,255,255,.08);background:#0d1116;position:sticky;top:0;z-index:5}'+
      '.dcr-module-list{display:flex;gap:14px}'+
      '.dcr-module-actions{margin-left:auto;display:flex;gap:8px;align-items:center}'+
      '.dcr-module-card{min-width:92px;height:42px;border:1px solid rgba(255,255,255,.14);background:#151b23;color:#e8eaed;border-radius:999px;padding:0 22px;font-size:18px;font-weight:900;display:inline-flex;align-items:center;justify-content:center;cursor:pointer}'+
      '.dcr-module-card.active{border-color:rgba(0,229,160,.65);background:#00e5a0;color:#00120c}'+
      '.dcr-child-toolbar{display:flex;align-items:center;gap:10px;padding:14px 22px 4px}'+
      '.dcr-source-switch{margin-left:auto;display:inline-flex;gap:6px;background:#0d1116;border:1px solid rgba(255,255,255,.1);border-radius:999px;padding:4px}'+
      '.dcr-src-btn{border:0;background:transparent;color:#8a9098;font-weight:800;font-size:13px;padding:6px 16px;border-radius:999px;cursor:pointer}'+
      '.dcr-src-btn.active{background:#00e5a0;color:#00120c}'+
      '.dcr-child-list{padding:8px 22px 24px;display:flex;flex-direction:column;gap:10px}'+
      '.dcr-child-card{border:1px solid rgba(255,255,255,.09);background:#111820;border-radius:14px;padding:12px;display:flex;gap:14px;align-items:center;position:relative}'+
      '.dcr-child-card.mpick{border-color:#00e5a0;background:rgba(0,229,160,.08)}'+
      '.dcr-child-card.mmode{cursor:pointer}'+
      '.dcr-thumb{width:120px;height:72px;flex-shrink:0;border-radius:10px;background:#050607;border:1px solid rgba(255,255,255,.08);overflow:hidden;display:flex;align-items:center;justify-content:center}'+
      '.dcr-thumb img{width:100%;height:100%;object-fit:contain}.dcr-thumb .noimg{color:#3a4452;font-size:11px}'+
      '.dcr-cmain{flex:1;min-width:0}.dcr-ctitle{font-size:15px;font-weight:900;margin-bottom:5px}'+
      '.dcr-cmeta{font-size:12px;color:#8a9098;line-height:1.7}.dcr-cmeta .k{color:#596271}'+
      '.dcr-cactions{display:flex;flex-direction:column;gap:6px;flex-shrink:0}'+
      '.dcr-badge{position:absolute;top:8px;left:8px;width:24px;height:24px;border-radius:50%;background:#00e5a0;color:#00120c;font-weight:900;display:flex;align-items:center;justify-content:center;font-size:13px;z-index:2}'+
      '.dcr-empty{text-align:center;color:#596271;padding:50px 10px;font-size:13px}'+
      '.dcr-preview{position:fixed;inset:0;background:rgba(0,0,0,.85);display:none;align-items:center;justify-content:center;z-index:3200}'+
      '.dcr-preview.show{display:flex}.dcr-preview img{max-width:92vw;max-height:92vh;object-fit:contain}';
    document.head.appendChild(st);
  }

  function render(){
    var ov=document.getElementById('dcReplicaOverlay'); if(!ov) return;
    document.querySelectorAll('#dcrSourceSwitch .dcr-src-btn').forEach(function(b){b.classList.toggle('active',b.getAttribute('data-src')===curSource)});
    var pl=document.getElementById('dcrParentList');
    pl.innerHTML=DATA.parents.length?DATA.parents.map(function(p){
      var c=pCounts(p);
      return '<div class="dcr-parent-card '+(p.parent_id===curParent?'active':'')+'" onclick="window.__dcReplica.selectParent(\''+E(p.parent_id)+'\')">'+
        '<div class="pn">'+E(p.parent_name)+'</div>'+
        '<div class="pc">美术 '+(c.material||0)+' · 分镜 '+(c.image||0)+' · 视频 '+(c.video||0)+'</div>'+
        '<div class="pt">最近更新：'+E(p.updated_at||'-')+'</div>'+
        (mMode?'':'<div><button class="btn small" onclick="event.stopPropagation();window.__dcReplica.dlParent(\''+E(p.parent_id)+'\',\''+E(p.parent_name)+'\')">下载父项目</button></div>')+
        '</div>';
    }).join(''):'<div class="dcr-desc">暂无项目。</div>';
    document.getElementById('dcrModuleList').innerHTML=DATA.module_order.map(function(mt){
      return '<button class="dcr-module-card '+(mt===curModule?'active':'')+'" onclick="window.__dcReplica.module(\''+mt+'\')">'+E(DATA.module_labels[mt]||mt)+'</button>';
    }).join('');
    var p=curParentItem();
    document.getElementById('dcrParentTitle').innerHTML='<b>当前父项目：</b>'+(p?E(p.parent_name):'未选择');
    var kids=p?(p.children||[]).filter(function(c){return c.module_type===curModule&&srcMatch(c)}):[];
    document.getElementById('dcrModuleCount').textContent=(DATA.module_labels[curModule]||curModule)+' '+kids.length;
    var cl=document.getElementById('dcrChildList');
    if(!p){cl.innerHTML='<div class="dcr-empty">请选择左侧父项目。</div>';return}
    cl.innerHTML=kids.length?kids.map(function(c){
      var gens=(c.generator_names&&c.generator_names.length)?c.generator_names.join('、'):(c.generator_name||'-');
      var thumb=c.has_preview?'<img src="/api/datacenter/preview/'+E(c.child_id)+'?t='+Date.now()+'" onerror="this.parentNode.innerHTML=\'<span class=&quot;noimg&quot;>无预览</span>\'">':'<span class="noimg">无预览</span>';
      var order=mSel.indexOf(c.child_id);
      var badge=(mMode&&order>=0)?'<span class="dcr-badge">'+(order+1)+'</span>':'';
      var click=mMode?'onclick="window.__dcReplica.pick(\''+E(c.child_id)+'\')"':'';
      var actions=mMode?'':'<div class="dcr-cactions">'+
        '<button class="btn small" onclick="window.__dcReplica.dlChild(\''+E(c.child_id)+'\',\''+E(c.child_name)+'\')">下载到本端</button>'+
        '<button class="btn ghost small" onclick="window.__dcReplica.preview(\''+E(c.child_id)+'\')">预览</button>'+'<button class="btn danger small" onclick="window.__dcReplica.del(\''+E(c.child_id)+'\',\''+E(c.child_name)+'\')">删除</button></div>';
      return '<div class="dcr-child-card '+(mMode?'mmode':'')+' '+(order>=0?'mpick':'')+'" '+click+'>'+badge+
        '<div class="dcr-thumb">'+thumb+'</div>'+
        '<div class="dcr-cmain"><div class="dcr-ctitle">'+E(c.child_name)+'</div>'+
        '<div class="dcr-cmeta"><span class="k">模块：</span>'+E(c.module_label)+'　<span class="k">生成者：</span>'+E(gens)+'<br>'+
        '<span class="k">来源：</span>'+E(c.source_label||'-')+'　<span class="k">上传时间：</span>'+E(c.uploaded_at||'-')+'　<span class="k">大小：</span>'+fmtSize(c.file_size)+'</div></div>'+
        actions+'</div>';
    }).join(''):'<div class="dcr-empty">该父项目在「'+(curSource==='reviewer'?'审核端':'制作端')+' · '+E(DATA.module_labels[curModule]||curModule)+'」下暂无子项目。</div>';
  }

  async function refresh(){
    try{
      var r=await dcApi('/api/datacenter/projects');
      DATA=Object.assign(DATA,r);
      if(!curParent&&r.parents.length)curParent=r.parents[0].parent_id;
      if(curParent&&!r.parents.some(function(p){return p.parent_id===curParent}))curParent=r.parents[0]?r.parents[0].parent_id:null;
      render();
    }catch(e){ document.getElementById('dcrChildList').innerHTML='<div class="dcr-empty">读取数据中心失败：'+E(e.message)+'<br>请检查数据中心地址、服务端是否启动、是否在同一局域网。</div>'; }
  }

  window.__dcReplica={
    open:function(hostLabel,opts){ ensureDom(); window.__dcrAllowMerge=!(opts&&opts.merge===false); document.getElementById('dcrHostTag').textContent=hostLabel||''; mMode=false;mSel=[];
      var mt=document.getElementById('dcrMergeToggle'); if(mt) mt.style.display=window.__dcrAllowMerge?'':'none';
      document.getElementById('dcrMergeToggle').textContent='合并子项目';document.getElementById('dcrMergeToggle').classList.add('ghost');document.getElementById('dcrMergeRun').style.display='none';
      document.getElementById('dcReplicaOverlay').classList.add('show'); refresh(); },
    close:function(){ var o=document.getElementById('dcReplicaOverlay'); if(o)o.classList.remove('show'); },
    refresh:refresh,
    selectParent:function(pid){curParent=pid;render()},
    module:function(mt){curModule=mt;render()},
    source:function(s){curSource=s;render()},
    preview:function(cid){var img=document.getElementById('dcrPreviewImg');img.src='/api/datacenter/preview/'+encodeURIComponent(cid)+'?t='+Date.now();img.onerror=function(){dcToast('该子项目没有可用预览图');document.getElementById('dcrPreview').classList.remove('show')};document.getElementById('dcrPreview').classList.add('show')},
    del:async function(cid,name){ if(!confirm('确定从数据中心删除子项目【'+name+'】？原始工程包会一并删除，且无法恢复。'))return; try{ await dcApi('/api/datacenter/delete-child',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({child_id:cid})}); dcToast('已从数据中心删除：'+name); await refresh(); }catch(e){ dcToast('删除失败：'+e.message); } },
    dlChild:function(cid,name){ if(typeof window.__dcHostDownloadChild==='function') window.__dcHostDownloadChild(cid,name); else dcToast('当前端未配置下载归位'); },
    dlParent:function(pid,name){ if(typeof window.__dcHostDownloadParent==='function') window.__dcHostDownloadParent(pid,name); else dcToast('当前端未配置下载归位'); },
    toggleMerge:function(){mMode=!mMode;mSel=[];var tb=document.getElementById('dcrMergeToggle'),rb=document.getElementById('dcrMergeRun');
      if(mMode){tb.textContent='取消合并';tb.classList.remove('ghost');rb.style.display='';rb.textContent='执行合并(0)'}else{tb.textContent='合并子项目';tb.classList.add('ghost');rb.style.display='none'}render()},
    pick:function(cid){var i=mSel.indexOf(cid);if(i>=0)mSel.splice(i,1);else mSel.push(cid);document.getElementById('dcrMergeRun').textContent='执行合并('+mSel.length+')';render()},
    runMerge:async function(){
      if(mSel.length<2)return dcToast('请按顺序选择至少 2 个子项目（同一父项目、同一模块）');
      var p=curParentItem(); var first=p?(p.children||[]).find(function(c){return c.child_id===mSel[0]}):null;
      var def=first?(first.base_child_name||first.child_name):'合并子项目';
      var name=prompt('合并后子项目名称（默认用第一个选中项目的名字）：',def||'合并子项目'); if(name===null)return;
      dcToast('正在数据中心合并 '+mSel.length+' 个子项目…');
      try{ await dcApi('/api/datacenter/merge',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({child_ids:mSel,target_name:(name||'').trim(),generator_name:(window.currentUserName||window.deviceNick||'')})});
        dcToast('合并完成，已存回数据中心'); mMode=false;mSel=[];
        document.getElementById('dcrMergeToggle').textContent='合并子项目';document.getElementById('dcrMergeToggle').classList.add('ghost');document.getElementById('dcrMergeRun').style.display='none';
        await refresh();
      }catch(e){ dcToast('合并失败：'+e.message); }
    }
  };
})();

window.__dcHostDownloadChild=function(cid,name){return dcDownloadChild(cid,name)};
window.__dcHostDownloadParent=function(pid,name){return dcDownloadParent(pid,name)};
window.dcOpenDownload=function(){window.__dcReplica.open('审核端 · 下载后自动归位到审核端',{merge:false})};

loadProjects();

</script></body></html>"""

class Handler(BaseHTTPRequestHandler):
    server_version='DilanUnifiedReview/1.0'
    def log_message(self, fmt, *args):
        sys.stderr.write('[%s] %s\n' % (self.log_date_time_string(), fmt%args))
    def send_bytes(self, status, data, ctype='application/octet-stream', headers=None):
        self.send_response(status)
        self.send_header('Content-Type',ctype)
        self.send_header('Content-Length',str(len(data)))
        if headers:
            for k,v in headers.items(): self.send_header(k,v)
        self.end_headers(); self.wfile.write(data)
    def send_json(self, status, data):
        self.send_bytes(status, json.dumps(data,ensure_ascii=False).encode('utf-8'), 'application/json; charset=utf-8')
    def read_body(self):
        n=int(self.headers.get('Content-Length') or 0)
        raw=self.rfile.read(n) if n else b'{}'
        return json.loads(raw.decode('utf-8') or '{}')
    def do_GET(self):
        u=urlparse(self.path); path=u.path; qs=parse_qs(u.query)
        try:
            if path=='/': self.send_bytes(200, INDEX_HTML.encode('utf-8'), 'text/html; charset=utf-8')
            elif path=='/api/projects': self.api_projects()
            elif path.startswith('/file/'):
                rest=path[len('/file/'):]
                wid,_,p=rest.partition('/')
                self.api_file(unquote(wid), unquote(p))
            elif path=='/api/workspace': self.api_workspace(qs.get('workspace',[''])[0])
            elif path=='/api/export_excel': self.api_export_excel(qs.get('workspace',[''])[0])
            elif path=='/api/export_package': self.api_export_package(qs.get('workspace',[''])[0])
            elif path=='/api/datacenter/config': self.api_datacenter_config_get()
            elif path=='/api/datacenter/projects': self.api_datacenter_projects()
            elif path.startswith('/api/datacenter/preview/'): self.api_datacenter_preview(path.rsplit('/',1)[-1])
            else: self.send_error(404)
        except Exception as e:
            self.send_bytes(500, str(e).encode('utf-8'), 'text/plain; charset=utf-8')
    def do_POST(self):
        u=urlparse(self.path)
        try:
            if u.path=='/api/import': self.api_import()
            elif u.path=='/api/discard': self.api_discard()
            elif u.path=='/api/delete_workspace': self.api_delete_workspace()
            elif u.path=='/api/create_parent': self.api_create_parent()
            elif u.path=='/api/rename_parent': self.api_rename_parent()
            elif u.path=='/api/rename_child': self.api_rename_child()
            elif u.path=='/api/delete_parent': self.api_delete_parent()
            elif u.path=='/api/review_submit': self.api_review_submit()
            elif u.path=='/api/review_save_times': self.api_review_save_times()
            elif u.path=='/api/review_add_note_image': self.api_review_add_note_image()
            elif u.path=='/api/review_add_note_image_url': self.api_review_add_note_image_url()
            elif u.path=='/api/review_save_note_images': self.api_review_save_note_images()
            elif u.path=='/api/review_delete_candidate': self.api_review_delete_candidate()
            elif u.path=='/api/datacenter/config/save': self.api_datacenter_config_save()
            elif u.path=='/api/datacenter/upload': self.api_datacenter_upload()
            elif u.path=='/api/datacenter/resolve-conflict': self.api_datacenter_resolve_conflict()
            elif u.path=='/api/datacenter/download-child': self.api_datacenter_download_child()
            elif u.path=='/api/datacenter/download-parent': self.api_datacenter_download_parent()
            elif u.path=='/api/datacenter/delete-child': self.api_datacenter_delete_child()
            else: self.send_error(404)
        except Exception as e:
            self.send_bytes(500, str(e).encode('utf-8'), 'text/plain; charset=utf-8')
    def api_projects(self):
        self.send_json(200, {'ok': True, 'parents': collect_review_parent_items(), 'projects': list_workspace_infos(), 'module_labels': MODULE_UI_LABELS})
    def api_create_parent(self):
        body=self.read_body(); name=body.get('parent_name') or body.get('name') or ''
        if not str(name).strip(): raise ValueError('请输入父项目名称')
        pid=safe_name(name,'父级项目')
        existing=find_review_parent(pid,name)
        if existing:
            parent=merge_review_parent(existing['parent_id'], existing.get('parent_name') or name)
        else:
            parent=merge_review_parent(pid, name)
        self.send_json(200, {'ok': True, 'parent': parent})
    def api_rename_child(self):
        body=self.read_body(); wid=body.get('workspace') or ''
        new_name=safe_name(body.get('child_name') or body.get('new_name') or '', '')
        if not wid: raise ValueError('missing workspace')
        if not new_name: raise ValueError('请输入新的子项目名称')
        wdir=workspace_dir(wid)
        info=read_json(wdir/'workspace.json',{}) or {}
        if not info: raise ValueError('找不到该子项目')
        info['child_name']=new_name; write_json(wdir/'workspace.json', info)
        # 同步进 project.json（重新上传时带上新名）
        try:
            w,inf2,project,manifest,usage=load_workspace(wid)
            project['child_name']=new_name
            scope=project.get('project_scope')
            if isinstance(scope,dict): scope['child_name']=new_name
            save_workspace_project(wid,project)
        except Exception:
            pass
        self.send_json(200, {'ok': True, 'child_name': new_name})

    def api_rename_parent(self):
        body=self.read_body(); pid=safe_name(body.get('parent_id') or '', '')
        new_name=safe_name(body.get('parent_name') or body.get('new_name') or '', '')
        if not pid: raise ValueError('missing parent_id')
        if not new_name: raise ValueError('请输入新的父项目名称')
        parent=merge_review_parent(pid, new_name, updated_at=now_str())
        if WORKSPACES.exists():
            for w in WORKSPACES.iterdir():
                if not w.is_dir(): continue
                info=read_json(w/'workspace.json',{}) or {}
                if safe_name(info.get('parent_id') or '', '') == pid:
                    info['parent_name']=new_name; write_json(w/'workspace.json', info)
        self.send_json(200, {'ok': True, 'parent': parent})
    def api_delete_parent(self):
        body=self.read_body(); pid=safe_name(body.get('parent_id') or '', '')
        if not pid: raise ValueError('missing parent_id')
        if WORKSPACES.exists():
            for w in list(WORKSPACES.iterdir()):
                if not w.is_dir(): continue
                info=read_json(w/'workspace.json',{}) or {}
                project=read_json(w/'project.json',{}) or {}
                manifest=read_json(w/'manifest.json',{}) or {}
                pinf=parent_info_from_workspace(info,project,manifest)
                if pinf.get('parent_id') == pid:
                    shutil.rmtree(w,ignore_errors=True)
        remove_review_parent(pid)
        self.send_json(200, {'ok': True})
    def api_delete_workspace(self):
        body=self.read_body(); wid=body.get('workspace') or ''
        w=workspace_dir(wid)
        try:
            if w.exists() and w.is_dir() and w.parent.resolve()==WORKSPACES.resolve():
                shutil.rmtree(w,ignore_errors=True)
        except Exception:
            pass
        self.send_json(200, {'ok': True})
    def api_file(self,wid,p):
        w=workspace_dir(wid)
        fp=resolve_file_path(w,p)
        if not fp or not fp.exists():
            self.send_error(404); return
        ctype=mimetypes.guess_type(fp.name)[0] or 'application/octet-stream'
        total=fp.stat().st_size
        rng=self.headers.get('Range') or self.headers.get('range')
        if rng and rng.strip().lower().startswith('bytes='):
            try:
                spec=rng.split('=',1)[1].split(',')[0].strip()
                start_s,end_s=(spec.split('-',1)+[''])[:2]
                if start_s=='':
                    # 末尾 N 字节
                    n=int(end_s); start=max(0,total-n); end=total-1
                else:
                    start=int(start_s); end=int(end_s) if end_s else total-1
                if start>=total: start=total-1
                if end>=total: end=total-1
                if end<start: end=start
                with fp.open('rb') as f:
                    f.seek(start); chunk=f.read(end-start+1)
                self.send_response(206)
                self.send_header('Content-Type',ctype)
                self.send_header('Content-Range','bytes %d-%d/%d'%(start,end,total))
                self.send_header('Accept-Ranges','bytes')
                self.send_header('Content-Length',str(len(chunk)))
                self.end_headers(); self.wfile.write(chunk); return
            except Exception:
                pass
        self.send_response(200)
        self.send_header('Content-Type',ctype)
        self.send_header('Accept-Ranges','bytes')
        self.send_header('Content-Length',str(total))
        self.end_headers(); self.wfile.write(fp.read_bytes())
    def api_import(self):
        filename,data=parse_multipart_upload(self.headers,self.rfile)
        self.send_json(200, import_review_zip(data, filename))
    def api_discard(self):
        body=self.read_body(); wid=body.get('workspace') or ''
        w=workspace_dir(wid)
        if w.exists() and w.is_dir() and w.parent==WORKSPACES:
            shutil.rmtree(w,ignore_errors=True)
        self.send_json(200,{'ok':True})
    def api_workspace(self,wid):
        if not wid: raise ValueError('missing workspace')
        w,info,project,manifest,usage=load_workspace(wid)
        if not project: raise FileNotFoundError('workspace not found')
        module_type=info.get('module_type') or detect_module(manifest,project,info.get('project_json_rel',''))
        if module_type=='material':
            ensure_material_review(project); project['__material_units']=material_units(project); save_workspace_project(wid,project)
        else:
            for sc in project.get('scenes') or []: ensure_scene_review(sc)
        self.send_json(200,{'ok':True,'workspace':wid,'info':info,'module_type':module_type,'project':project,'manifest':manifest,'usage_summary':usage})
    def api_review_save_note_images(self):
        body=self.read_body(); wid=body.get('workspace') or ''
        unit_type=body.get('unit_type') or 'scene'; uid=body.get('unit_id') or ''
        round_no=str(body.get('round') or 1); images=body.get('images') or []
        w,info,project,manifest,usage=load_workspace(wid)
        all_scenes=[]
        wss=project.get('material_workspaces') if isinstance(project.get('material_workspaces'),dict) else {}
        for key in ('character','scene','object'):
            for s in (wss.get(key) or {}).get('scenes') or []: all_scenes.append(s)
        for s in project.get('scenes') or []: all_scenes.append(s)
        target=None
        if unit_type=='shot':
            for s in all_scenes:
                for sh in (s.get('shots') or []):
                    if sh.get('shot_id')==uid: target=sh; break
                if target: break
        else:
            for s in all_scenes:
                if (s.get('scene_id') or scene_code(s))==uid or scene_code(s)==uid: target=s; break
        if target is None: self.send_json(404,{'error':'找不到审核单位'}); return
        nbr=target.setdefault('review_notes_by_round',{})
        slot=nbr.setdefault(round_no,{'text':'','images':[]})
        slot['images']=images
        save_workspace_project(wid,project)
        self.send_json(200,{'ok':True})

    def api_review_delete_candidate(self):
        body=self.read_body(); wid=body.get('workspace') or ''
        uid=body.get('shot_id') or ''
        cand_id=body.get('candidate_id'); idx=body.get('index')
        w,info,project,manifest,usage=load_workspace(wid)
        all_scenes=[]
        wss=project.get('material_workspaces') if isinstance(project.get('material_workspaces'),dict) else {}
        for key in ('character','scene','object'):
            for s in (wss.get(key) or {}).get('scenes') or []: all_scenes.append(s)
        for s in project.get('scenes') or []: all_scenes.append(s)
        target=None
        for s in all_scenes:
            for sh in (s.get('shots') or []):
                if sh.get('shot_id')==uid: target=sh; break
            if target: break
        if target is None: self.send_json(404,{'error':'找不到子栏'}); return
        cands=target.get('storyboard_candidates') or []
        visible=[c for c in cands if not c.get('review_deleted')]
        chosen=None
        if cand_id not in (None,''):
            for c in cands:
                if str(c.get('image_id') or c.get('candidate_id') or c.get('id') or '')==str(cand_id): chosen=c; break
        if chosen is None and isinstance(idx,int) and 0<=idx<len(visible):
            chosen=visible[idx]
        if chosen is None: self.send_json(404,{'error':'找不到待选图'}); return
        chosen['review_deleted']=True
        save_workspace_project(wid,project)
        self.send_json(200,{'ok':True})

    def api_review_add_note_image_url(self):
        from urllib.parse import unquote
        body=self.read_body(); wid=body.get('workspace') or ''
        unit_type=body.get('unit_type') or 'scene'; uid=body.get('unit_id') or ''
        url=body.get('url') or ''; round_no=str(body.get('round') or 1)
        # 解析 /file/<wid>/<rel> 形式的本端文件
        rel=None
        if '/file/' in url:
            tail=url.split('/file/',1)[1]
            parts=tail.split('/',1)
            if len(parts)==2: rel=unquote(parts[1].split('?')[0].split('#')[0])
        elif url and not url.startswith(('http://','https://','data:')):
            rel=unquote(url.lstrip('/').split('?')[0])
        if not rel: self.send_json(400,{'error':'无法解析素材地址'}); return
        w,info,project,manifest,usage=load_workspace(wid)
        srcfp=resolve_file_path(w,rel)
        if not srcfp or not srcfp.exists(): self.send_json(404,{'error':'素材文件不存在'}); return
        ext=srcfp.suffix or '.png'
        # 定位 scene/shot
        all_scenes=[]
        wss=project.get('material_workspaces') if isinstance(project.get('material_workspaces'),dict) else {}
        for key in ('character','scene','object'):
            for s in (wss.get(key) or {}).get('scenes') or []: all_scenes.append(s)
        for s in project.get('scenes') or []: all_scenes.append(s)
        target=None
        if unit_type=='shot':
            for s in all_scenes:
                for sh in (s.get('shots') or []):
                    if sh.get('shot_id')==uid: target=sh; break
                if target: break
        else:
            for s in all_scenes:
                if (s.get('scene_id') or scene_code(s))==uid or scene_code(s)==uid: target=s; break
        if target is None: self.send_json(404,{'error':'找不到审核单位'}); return
        note_text=body.get('note_text')
        if note_text is not None:
            _rv=target.setdefault('review',{}); _rds=_rv.setdefault('rounds',{}); _rds.setdefault(round_no,{})['note']=note_text
        rel_dir=w/'review_notes'; rel_dir.mkdir(parents=True, exist_ok=True)
        nbr=target.setdefault('review_notes_by_round',{})
        slot=nbr.setdefault(round_no,{'text':'','images':[]})
        idx=len(slot.get('images') or [])
        fname=safe_name(f"note_{uid}_{round_no}_{idx}")+ext
        (rel_dir/fname).write_bytes(srcfp.read_bytes())
        relout='review_notes/'+fname
        slot.setdefault('images',[]).append({'file_path':relout,'added_at':now_str()})
        save_workspace_project(wid,project)
        self.send_json(200,{'ok':True,'project':project,'file_path':relout})

    def api_review_add_note_image(self):
        import base64 as _b64
        body=self.read_body(); wid=body.get('workspace') or ''
        unit_type=body.get('unit_type') or 'scene'; uid=body.get('unit_id') or ''
        data_url=body.get('dataUrl') or ''; round_no=str(body.get('round') or 1)
        if ',' not in data_url: self.send_json(400,{'error':'missing image'}); return
        head,b64=data_url.split(',',1); raw=_b64.b64decode(b64)
        ext='.png'
        if 'jpeg' in head or 'jpg' in head: ext='.jpg'
        elif 'webp' in head: ext='.webp'
        elif 'gif' in head: ext='.gif'
        w,info,project,manifest,usage=load_workspace(wid)
        # 定位 scene 或 shot（含美术工作台）
        all_scenes=[]
        wss=project.get('material_workspaces') if isinstance(project.get('material_workspaces'),dict) else {}
        for key in ('character','scene','object'):
            for s in (wss.get(key) or {}).get('scenes') or []: all_scenes.append(s)
        for s in project.get('scenes') or []: all_scenes.append(s)
        target=None
        if unit_type=='shot':
            for s in all_scenes:
                for sh in (s.get('shots') or []):
                    if sh.get('shot_id')==uid: target=sh; break
                if target: break
        else:
            for s in all_scenes:
                if (s.get('scene_id') or scene_code(s))==uid or scene_code(s)==uid: target=s; break
        if target is None: self.send_json(404,{'error':'找不到审核单位'}); return
        note_text=body.get('note_text')
        if note_text is not None:
            _rv=target.setdefault('review',{}); _rds=_rv.setdefault('rounds',{}); _rds.setdefault(round_no,{})['note']=note_text
        rel_dir=w/'review_notes'; rel_dir.mkdir(parents=True, exist_ok=True)
        nbr=target.setdefault('review_notes_by_round',{})
        slot=nbr.setdefault(round_no,{'text':'','images':[]})
        idx=len(slot.get('images') or [])
        fname=safe_name(f"note_{uid}_{round_no}_{idx}")+ext
        (rel_dir/fname).write_bytes(raw)
        rel='review_notes/'+fname
        slot.setdefault('images',[]).append({'file_path':rel,'added_at':now_str()})
        save_workspace_project(wid,project)
        self.send_json(200,{'ok':True,'project':project,'file_path':rel})

    def api_review_save_times(self):
        body=self.read_body(); wid=body.get('workspace') or ''; times=body.get('times') or []
        w,info,project,manifest,usage=load_workspace(wid)
        # 建立 scene_id / shot_id -> 对象 的索引（含美术三套 workspace + 普通 scenes）
        all_scenes=[]
        wss=project.get('material_workspaces') if isinstance(project.get('material_workspaces'),dict) else {}
        for key in ('character','scene','object'):
            for s in (wss.get(key) or {}).get('scenes') or []:
                all_scenes.append(s)
        for s in project.get('scenes') or []:
            all_scenes.append(s)
        def sid_of(s):
            return s.get('scene_id') or scene_code(s)
        scene_idx={sid_of(s):s for s in all_scenes}
        shot_idx={}
        for s in all_scenes:
            for sh in (s.get('shots') or []):
                if sh.get('shot_id'): shot_idx[sh['shot_id']]=sh
        for t in times:
            sc=scene_idx.get(t.get('scene_id'))
            if sc is not None:
                if t.get('time_start') is not None: sc['time_start']=t['time_start']
                if t.get('time_end') is not None: sc['time_end']=t['time_end']
            for st in (t.get('shots') or []):
                sh=shot_idx.get(st.get('shot_id'))
                if sh is not None and st.get('duration_seconds') is not None:
                    sh['duration_seconds']=st['duration_seconds']
        save_workspace_project(wid,project)
        self.send_json(200,{'ok':True})

    def api_review_submit(self):
        body=self.read_body(); wid=body.get('workspace') or ''; status=body.get('status') or 'pending'; note=body.get('note') or ''; reviewer=body.get('reviewer') or '审核员'
        w,info,project,manifest,usage=load_workspace(wid)
        module_type=info.get('module_type') or body.get('module_type') or detect_module(manifest,project,info.get('project_json_rel',''))
        unit_type=body.get('unit_type') or ''
        uid=body.get('unit_id') or ''
        if module_type=='material' and unit_type in ('scene','shot'):
            # 美术工作台审核：在三套 workspace 的 scenes 里定位命名框(scene)或道具子栏(shot)
            target=None
            wss=project.get('material_workspaces') if isinstance(project.get('material_workspaces'),dict) else {}
            search_scopes=[]
            for key in ('character','scene','object'):
                ws=wss.get(key) or {}
                search_scopes.append(ws.get('scenes') or [])
            search_scopes.append(project.get('scenes') or [])
            if unit_type=='scene':
                for scenes in search_scopes:
                    for s in scenes:
                        if (s.get('scene_id') or scene_code(s))==uid or scene_code(s)==uid:
                            target=s; break
                    if target: break
                if target is None: raise ValueError('找不到审核单位')
                rv=ensure_scene_review(target)
                update_round(rv, rv.get('current_round') or 1, status, reviewer, note)
            else:  # shot（道具单层）
                for scenes in search_scopes:
                    for s in scenes:
                        for sh in (s.get('shots') or []):
                            if sh.get('shot_id')==uid:
                                target=sh; break
                        if target: break
                    if target: break
                if target is None: raise ValueError('找不到审核单位')
                rv=target.setdefault('review',{}); ensure_rounds(rv)
                update_round(rv, rv.get('current_round') or 1, status, reviewer, note)
            save_workspace_project(wid,project)
            self.send_json(200,{'ok':True,'project':project}); return
        if body.get('unit_type')=='material' or module_type=='material':
            ensure_material_review(project); uid=body.get('unit_id') or ''
            rv=project['material_review']['units'].setdefault(uid, {'current_round':1,'max_round_available':1,'rounds':{}})
            update_round(rv, rv.get('current_round') or 1, status, reviewer, note)
            project['__material_units']=material_units(project)
        else:
            uid=body.get('unit_id') or ''
            sc=None
            for s in project.get('scenes') or []:
                if (s.get('scene_id') or scene_code(s))==uid or scene_code(s)==uid:
                    sc=s; break
            if not sc: raise ValueError('找不到审核单位')
            rv=ensure_scene_review(sc)
            update_round(rv, rv.get('current_round') or 1, status, reviewer, note)
        save_workspace_project(wid,project)
        self.send_json(200,{'ok':True,'project':project})
    def api_export_excel(self,wid):
        w,info,project,manifest,usage=load_workspace(wid)
        module_type=info.get('module_type') or detect_module(manifest,project,info.get('project_json_rel',''))
        if module_type=='material': ensure_material_review(project)
        data=export_excel_bytes(project,usage,module_type,manifest)
        name=safe_name((project.get('project_name') or project.get('project_id') or '工程')+'_审核表_'+stamp())+'.xlsx'
        self.send_bytes(200,data,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',{'Content-Disposition':"attachment; filename*=UTF-8''"+quote(name)})
    def api_export_package(self,wid):
        w,info,project,manifest,usage=load_workspace(wid)
        data=build_reviewed_package(wid)
        name=safe_name((project.get('project_name') or project.get('project_id') or '工程')+'_reviewed_'+stamp())+'.zip'
        self.send_bytes(200,data,'application/zip',{'Content-Disposition':"attachment; filename*=UTF-8''"+quote(name)})

    # ---- 数据中心接入（V7） ----
    def _dc_dir(self):
        return ROOT
    def api_datacenter_config_get(self):
        self.send_json(200, dcc.load_dc_config(self._dc_dir()))
    def api_datacenter_config_save(self):
        body=self.read_body()
        self.send_json(200, dcc.save_dc_config(self._dc_dir(), body.get('url') or ''))
    def api_datacenter_projects(self):
        try: self.send_json(200, dcc.list_projects(self._dc_dir()))
        except Exception as e: self.send_json(502, {'error': str(e)})
    def api_datacenter_upload(self):
        body=self.read_body(); wid=body.get('workspace') or ''
        user_name=(body.get('user_name') or '审核员').strip() or '审核员'
        on_conflict=body.get('on_conflict') or 'ask'
        if not wid: self.send_json(400,{'error':'请先打开一个审核工程'}); return
        try:
            w,info,project,manifest,usage=load_workspace(wid)
            module_type=info.get('module_type') or detect_module(manifest,project,info.get('project_json_rel',''))
            data=build_reviewed_package(wid)
            pinf=parent_info_from_workspace(info, project, manifest)
            gen=generator_display(project, manifest) or user_name
            review_summary=make_review_result(project, manifest, usage, module_type)
            meta={'parent_id':pinf['parent_id'],'parent_name':pinf['parent_name'],
                  'module_type':module_type,'module_label':MODULE_LABELS.get(module_type,module_type),
                  'child_id':pinf['child_id'],'child_name':pinf['child_name'],
                  'base_child_name':pinf['child_name'],'generator_name':gen,
                  'client_type':'reviewer','review_status':review_summary}
            fname=safe_name((project.get('project_name') or 'reviewed'))+'.zip'
            self.send_json(200, dcc.upload_package(self._dc_dir(), data, fname, meta, on_conflict=on_conflict))
        except Exception as e:
            self.send_json(502, {'error': str(e)})
    def api_datacenter_resolve_conflict(self):
        body=self.read_body()
        try: self.send_json(200, dcc.resolve_conflict(self._dc_dir(), body.get('token') or '', body.get('action') or 'new_version'))
        except Exception as e: self.send_json(502, {'error': str(e)})
    def api_datacenter_download_child(self):
        body=self.read_body(); cid=body.get('child_id') or ''
        if not cid: self.send_json(400,{'error':'missing child_id'}); return
        try:
            raw=dcc.download_child_bytes(self._dc_dir(), cid)
            self.send_json(200, {'ok':True, **import_review_zip(raw, (body.get('name') or 'download')+'.zip')})
        except Exception as e:
            self.send_json(502, {'error': str(e)})
    def api_datacenter_download_parent(self):
        body=self.read_body(); pid=body.get('parent_id') or ''
        if not pid: self.send_json(400,{'error':'missing parent_id'}); return
        try:
            import base64 as _b64
            bundle=dcc.download_parent_children(self._dc_dir(), pid); ok=0
            kids=bundle.get('children') or []
            for k in kids:
                du=k['dataUrl']; b=_b64.b64decode(du.split(',',1)[1] if ',' in du else du)
                try: import_review_zip(b, k.get('filename') or 'child.zip'); ok+=1
                except Exception: pass
            self.send_json(200, {'ok':True,'parent_name':bundle.get('parent_name'),'imported':ok,'total':len(kids)})
        except Exception as e:
            self.send_json(502, {'error': str(e)})


    def api_datacenter_preview(self, child_id):
        try:
            raw=dcc.preview_bytes(self._dc_dir(), child_id)
            self.send_response(200); self.send_header('Content-Type','image/png')
            self.send_header('Content-Length',str(len(raw))); self.send_header('Cache-Control','no-store'); self.end_headers()
            self.wfile.write(raw)
        except Exception as e:
            self.send_json(404,{'error':str(e)})
    def api_datacenter_delete_child(self):
        body=self.read_body(); cid=body.get('child_id') or ''
        if not cid: self.send_json(400,{'error':'missing child_id'}); return
        try: self.send_json(200, dcc.delete_child(self._dc_dir(), cid))
        except Exception as e: self.send_json(502,{'error':str(e)})


def run():
    httpd=ThreadingHTTPServer(('127.0.0.1',PORT),Handler)
    url=f'http://127.0.0.1:{PORT}/'
    print('帝蓝工作流-统一审核端 V37 已启动：',url)
    threading.Timer(0.8,lambda:webbrowser.open(url)).start()
    httpd.serve_forever()

if __name__=='__main__':
    run()